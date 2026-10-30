"""Layout/order regressions using synthetic fixtures, not experimental evidence."""
import copy,unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
import test_biochemical as fixtures
from bio_common import read,write
from bio_author import build,replace_text,format_p
from bio_layout import table_format,table_evidence,normalize_plan,summary_format,analysis_evidence,visible,table_width_plan,section_content_width
import inspect_docx,bio_contract,batch_tool

class LayoutTests(unittest.TestCase):
    def fixture(self):
        f=fixtures.BiochemicalBehaviorTests();f.setUp();self.addCleanup(f.doCleanups);return f
    def test_source_table_values_units_scripts_and_merge_survive_dedicated_layout(self):
        d=Document();t=d.add_table(rows=3,cols=3);t.cell(0,0).merge(t.cell(0,1));t.cell(0,0).text='表头'
        p=t.cell(0,2).paragraphs[0];p.add_run('密度(g/cm');p.add_run('3').font.superscript=True;p.add_run(')')
        t.cell(1,0).text='样品A';t.cell(1,2).text='2.508±0.014'
        before=visible(t._tbl);grid=ET.tostring(t._tbl.tblGrid)
        for row in t.rows:
            rp=row._tr.get_or_add_trPr();height=OxmlElement('w:trHeight');height.set(qn('w:hRule'),'exact');rp.append(height)
        for p in t._tbl.xpath('.//w:p'):format_p(p,'table')
        table_format(t._tbl)
        self.assertEqual(visible(t._tbl),before);self.assertEqual(ET.tostring(t._tbl.tblGrid),grid)
        self.assertEqual(len(t._tbl.xpath('.//w:vertAlign')),1);self.assertFalse(table_evidence(d.element,'sha')['issues'])
        self.assertEqual(table_evidence(d.element,'sha')['tables'][0]['header_rows'],[0])
    def test_each_table_repeats_its_own_source_header_and_supports_multirow_header(self):
        d=Document();a=d.add_table(rows=4,cols=2);a.cell(0,0).text='力学性能';a.cell(1,0).text='抗弯强度(MPa)'
        b=d.add_table(rows=3,cols=3);b.cell(0,0).text='磨损量'
        table_format(a._tbl,2);table_format(b._tbl,1)
        ev=table_evidence(d.element,'sha');self.assertFalse(ev['issues'])
        self.assertEqual([t['header_rows'] for t in ev['tables']],[[0,1],[0]])
        self.assertEqual(ev['tables'][1]['header_texts'][0][0],'磨损量')
        b.rows[1]._tr.get_or_add_trPr().find(qn('w:tblHeader')).set(qn('w:val'),'1')
        b.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')).set(qn('w:val'),'0')
        self.assertIn('own_contiguous_repeated_header_missing',[i['code'] for i in table_evidence(d.element,'sha')['issues']])
    def test_inline_result_heading_split_and_all_figure_analyses_go_first(self):
        f=self.fixture();p=f.plan();original=copy.deepcopy(p);normalized,changes=normalize_plan(p,f.e)
        self.assertEqual(p,original)
        out=f.root/'draft.docx';write(f.root/'plan.json',p);build(f.root/'evidence/evidence.json',f.root/'plan.json',out)
        texts=[p.text for p in Document(out).paragraphs];idx=texts.index('结果分析：')
        self.assertEqual(texts[idx+1:idx+4],['由图1可知，示例标签保留。','由图2可知，示例标签保留。','这是测试夹具。'])
        self.assertEqual(next(x['count'] for x in changes if x['kind']=='figure_analyses_first_in_results'),2)
    def test_summary_removes_label_only_and_adds_subject_without_replacing_facts(self):
        d=Document();p=d.add_paragraph('技术总结：以材料A承担主体载荷，保留既有结论。')
        info=summary_format(p._p,{},replace_text)
        self.assertEqual(p.text,'本发明以材料A承担主体载荷，保留既有结论。');self.assertFalse(info['technical_content_deleted'])
        p=d.add_paragraph('技术总结：本发明采用原方法。');summary_format(p._p,{},replace_text)
        self.assertEqual(p.text,'本发明采用原方法。')
        p=d.add_paragraph('技术总结：该材料保持原状态。');summary_format(p._p,{},replace_text)
        self.assertEqual(p.text,'该材料保持原状态。')
        p=d.add_paragraph('技术总结性评价仍需保留原始条件。');summary_format(p._p,{},replace_text)
        self.assertEqual(p.text,'技术总结性评价仍需保留原始条件。')
    def test_separate_summary_label_is_removed_and_next_paragraph_preserved(self):
        e={'body':[{'body_index':0,'text':'技术总结：'},{'body_index':1,'text':'采用原工序。'}]}
        p={'sections':{'description':[{'source_body_index':0},{'source_body_index':1}]}}
        for row in e['body']:row.setdefault('kind','p')
        norm,_=normalize_plan(p,e);self.assertEqual(len(norm['sections']['description']),1)
        self.assertEqual(norm['sections']['description'][0]['role'],'technical_summary')
    def test_multiline_figure_analysis_and_end_figure_captions_do_not_false_fail(self):
        rows=[{'p':i,'text':t,'body_direct':True,'section_headers':[section],'biochemical_role':role} for i,(t,section,role) in enumerate([
            ('结果分析：','说明书','BioResultsHeading'),('由图1可知，保留来源。','说明书','BioFigureAnalysis'),('此处是同一图的第二段。','说明书','BioFigureAnalysis'),
            ('原综合结果。','说明书',None),('图1','说明书附图',None)])]
        self.assertFalse(analysis_evidence(rows,'sha')['issues'])
        rows.insert(1,dict(rows[3],p=0.5));self.assertIn('figure_analysis_not_first_in_results',[i['code'] for i in analysis_evidence(rows,'sha')['issues']])
    def test_actual_inspection_blocks_reintroduced_indent_and_summary_label(self):
        f=self.fixture();out=f.build();d=Document(out);p=d.tables[0].cell(0,0).paragraphs[0]
        p.paragraph_format.first_line_indent=120000
        result=next(p for p in d.paragraphs if p.text=='结果分析：');result.insert_paragraph_before('技术总结：以原材料处理。')
        tampered=f.root/'tampered.docx';d.save(tampered)
        s=inspect_docx.run_inspection(tampered,f.root/'inspection','draft',True,True);report=read(s['inspection'])
        self.assertTrue(report['biochemical_table_evidence']['issues']);self.assertTrue(report['biochemical_analysis_evidence']['issues'])
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(report,f.e['title_candidate'])
    def test_pass_requires_current_complete_layout_evidence(self):
        f=self.fixture();out=f.build();s=inspect_docx.run_inspection(out,f.root/'inspection','draft',True,True);report=read(s['inspection'])
        bio_contract.check(report,f.e['title_candidate'])
        for key in ('biochemical_table_evidence','biochemical_analysis_evidence','biochemical_presentation_evidence'):
            for mutation in ('missing','stale','incomplete'):
                bad=copy.deepcopy(report)
                if mutation=='missing':bad.pop(key)
                elif mutation=='stale':bad[key]['artifact_sha256']='0'*64
                else:bad[key]['complete']=False
                with self.assertRaises(batch_tool.BatchError):bio_contract.check(bad,f.e['title_candidate'])

    def text_table(self,values):
        d=Document();t=d.add_table(rows=len(values),cols=len(values[0]))
        for row,strings in zip(t.rows,values):
            for cell,value in zip(row.cells,strings):cell.text=value
        for p in t._tbl.xpath('.//w:p'):format_p(p,'table')
        return d,t

    def test_short_identifier_yields_width_to_long_description_without_content_changes(self):
        d,t=self.text_table([['样品','观察说明'],['实施例123','吸收区域清晰，主峰位置保持一致，谱带连续且无异常。']])
        before=[ET.tostring(p) for p in t._tbl.xpath('.//w:p/w:r')]
        available=section_content_width(d.sections[0]._sectPr)
        result=table_format(t._tbl,available_twips=available)['width_plan']
        a,b=result['column_widths_twips']
        self.assertLess(a,b/2);self.assertEqual(a+b,available)
        self.assertFalse(result['estimated_wrapping_columns'])
        self.assertEqual(before,[ET.tostring(p) for p in t._tbl.xpath('.//w:p/w:r')])
        self.assertFalse(table_evidence(d.element,'sha')['issues'])

    def test_numeric_and_native_script_columns_keep_units_and_readable_size(self):
        d,t=self.text_table([['样品','指标A','单位'],['对比例12','1.234±0.005','cm']])
        run=t.cell(1,2).paragraphs[0].add_run('3');run.font.superscript=True
        for p in t._tbl.xpath('.//w:p'):format_p(p,'table')
        scripts=ET.tostring(run._r);data=visible(t._tbl)
        plan=table_format(t._tbl,available_twips=8640)['width_plan']
        self.assertFalse(plan['estimated_wrapping_columns'])
        self.assertEqual(visible(t._tbl),data);self.assertEqual(ET.tostring(run._r),scripts)
        self.assertTrue(all(r.get(qn('w:val'))=='24' for r in t._tbl.xpath('.//w:sz')))

    def test_content_that_cannot_fit_still_wraps_without_shrinking_or_no_wrap(self):
        d,t=self.text_table([['样品','观察说明'],['实施例1','保留完整文字。'*30]])
        pr=t.cell(1,1)._tc.get_or_add_tcPr()
        pr.append(OxmlElement('w:noWrap'));pr.append(OxmlElement('w:tcFitText'))
        original=visible(t._tbl)
        result=table_format(t._tbl,available_twips=8640)['width_plan']
        self.assertEqual(result['estimated_wrapping_columns'],[1])
        self.assertEqual(sum(result['column_widths_twips']),8640)
        self.assertEqual(visible(t._tbl),original)
        self.assertFalse(table_evidence(d.element,'sha')['issues'])

    def test_manual_line_break_and_distinct_paragraphs_survive(self):
        d,t=self.text_table([['样品','说明'],['实施例1','第一项']])
        t.cell(1,1).paragraphs[0].add_run().add_break()
        t.cell(1,1).paragraphs[0].add_run('第二项')
        t.cell(1,1).add_paragraph('独立说明')
        table_format(t._tbl,available_twips=8640)
        self.assertEqual(len(t._tbl.xpath('.//w:br')),1)
        self.assertEqual(len(t.cell(1,1).paragraphs),2)

    def test_rich_and_merged_tables_do_not_get_guessed_widths(self):
        d,t=self.text_table([['合并表头',''],['A','B']]);t.cell(0,0).merge(t.cell(0,1))
        original=ET.tostring(t._tbl.tblGrid)
        self.assertEqual(table_format(t._tbl,available_twips=8640)['width_plan']['mode'],'source_geometry_requires_review')
        self.assertEqual(ET.tostring(t._tbl.tblGrid),original)
        d,t=self.text_table([['样品','公式'],['A','']])
        t.cell(1,1).paragraphs[0]._p.append(OxmlElement('m:oMath'))
        original=ET.tostring(t._tbl.tblGrid)
        self.assertEqual(table_width_plan(t._tbl,8640)['reason'],'rich_or_generated_content')
        table_format(t._tbl,available_twips=8640)
        self.assertEqual(ET.tostring(t._tbl.tblGrid),original)

    def test_reformat_is_idempotent_and_keeps_repeated_header(self):
        d,t=self.text_table([['样品','长说明'],['A','保持所有文字与边框结构。']])
        table_format(t._tbl,available_twips=8640);original=ET.tostring(t._tbl)
        table_format(t._tbl,available_twips=8640)
        self.assertEqual(ET.tostring(t._tbl),original)
        self.assertEqual(table_evidence(d.element,'sha')['tables'][0]['header_rows'],[0])

    def test_inspection_detects_grid_drift_overflow_and_forced_fit(self):
        d,t=self.text_table([['样品','说明'],['A','普通文字']]);table_format(t._tbl,available_twips=8640)
        t.cell(1,1)._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),'400')
        t.cell(1,0)._tc.get_or_add_tcPr().find(qn('w:noWrap')).set(qn('w:val'),'1')
        t._tbl.tblGrid[0].set(qn('w:w'),'99999')
        ev=table_evidence(d.element,'sha')
        self.assertTrue({'table_exceeds_section_width','cell_grid_width_mismatch','forced_cell_text_fit'} <= {i['code'] for i in ev['issues']})
        # Both the preflight (lxml) and inspector (stdlib XML) paths must agree.
        import xml.etree.ElementTree as StdET
        self.assertEqual(ev,table_evidence(StdET.fromstring(ET.tostring(d.element)),'sha'))

    def test_width_contract_required_and_generator_records_actual_section_width(self):
        f=self.fixture();out=f.build()
        trace=read(out.with_suffix('.trace.json'))
        layouts=[b['table_layout']['width_plan'] for b in trace['blocks'] if 'table_layout' in b]
        self.assertTrue(layouts)
        width=section_content_width(Document(out).sections[3]._sectPr)
        self.assertTrue(all(p['available_width_twips']==width for p in layouts if p['mode']=='content_balanced'))
        report=read(inspect_docx.run_inspection(out,f.root/'inspection','draft',True,True)['inspection'])
        bio_contract.check(report,f.e['title_candidate'])
        report['biochemical_table_evidence'].pop('width_contract_version')
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(report,f.e['title_candidate'])
if __name__=='__main__':unittest.main()
