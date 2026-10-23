"""Native script behavior with the user's chemical and negative-exponent cases."""
from pathlib import Path
from copy import deepcopy
import tempfile,unittest,zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
from bio_scripts import normalize,scan_docx
from bio_common import read,write,sha
from bio_prepare import prepare
from bio_author import build
import test_biochemical as fixtures
import inspect_docx,bio_contract,batch_tool

def spans(p):
    return [(r.text,r.font.subscript,r.font.superscript) for r in p.runs if r.text]

class NativeScriptTests(unittest.TestCase):
    def fixture(self):
        f=fixtures.BiochemicalBehaviorTests();f.setUp();self.addCleanup(f.doCleanups);return f
    def test_chemical_subscripts_and_whole_negative_exponent_use_plain_native_runs(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'format.docx';d=Document();p=d.add_paragraph();r=p.add_run('Na₂CO₃');r.bold=True
            t=d.add_table(rows=1,cols=1);t.cell(0,0).text='10⁻⁵mm³/(N·m)';d.sections[0].header.paragraphs[0].text='cm²'
            d.add_paragraph('S1、实施例2、1000℃、β、γ、±、№3');d.save(path)
            normalize(path);d=Document(path)
            self.assertEqual(spans(d.paragraphs[0]),[('Na',None,None),('2',True,False),('CO',None,None),('3',True,False)])
            self.assertTrue(all(r.bold for r in d.paragraphs[0].runs))
            self.assertEqual(spans(d.tables[0].cell(0,0).paragraphs[0]),[('10',None,None),('-5',False,True),('mm',None,None),('3',False,True),('/(N·m)',None,None)])
            self.assertEqual(spans(d.sections[0].header.paragraphs[0]),[('cm',None,None),('2',False,True)])
            self.assertEqual(d.paragraphs[1].text,'S1、实施例2、1000℃、β、γ、±、№3')
            self.assertFalse(scan_docx(path)['issues']);before=sha(path);normalize(path);self.assertEqual(sha(path),before)
    def test_existing_native_scripts_and_omml_and_images_are_unchanged(self):
        f=self.fixture();out=f.build()
        with zipfile.ZipFile(out) as z:before={n:z.read(n) for n in z.namelist()}
        result=normalize(out)
        with zipfile.ZipFile(out) as z:after={n:z.read(n) for n in z.namelist()}
        self.assertFalse(result['converted_runs']);self.assertEqual(before,after)
    def test_complex_math_or_nested_conflicting_scripts_fail_without_partial_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            for kind in ('math','opposite','mixed_object'):
                d=Document();d.add_paragraph('Na₂CO₃');p=d.add_paragraph()
                if kind=='math':
                    math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x²';r.append(t);math.append(r);p._p.append(math)
                else:
                    r=p.add_run('a₂')
                    if kind=='opposite':r.font.superscript=True
                    else:r.add_break()
                path=Path(tmp)/(kind+'.docx');d.save(path);before=sha(path)
                with self.assertRaises(ValueError):normalize(path)
                self.assertEqual(sha(path),before)
    def test_source_glyphs_in_body_and_table_reach_native_format_through_generator(self):
        f=self.fixture();d=Document(f.source);p=next(p for p in d.paragraphs if p.text.startswith('保留(SO'))
        p.add_run('Na₂CO₃');d.tables[0].cell(0,1).text='10⁻⁵mm³/(N·m)';d.save(f.source)
        before=sha(f.source);e=prepare(f.source,f.root/'native-source');f.e=e
        self.assertTrue(any(r['unicode_script_glyphs'] for r in e['body']))
        plan=f.plan();write(f.root/'native-plan.json',plan);out=f.root/'native-draft.docx';build(f.root/'native-source/evidence.json',f.root/'native-plan.json',out)
        dest=Document(out);self.assertEqual(sha(f.source),before);self.assertFalse(scan_docx(out)['issues'])
        p=dest.tables[0].cell(0,1).paragraphs[0]
        self.assertTrue(any(r.text=='-5' and r.font.superscript for r in p.runs))
        chemical=next(p for p in dest.paragraphs if 'Na2CO3' in p.text)
        self.assertTrue(any(r.text=='2' and r.font.subscript for r in chemical.runs))
        self.assertTrue(any(r.text=='3' and r.font.subscript for r in chemical.runs))
    def test_final_gate_blocks_reintroduced_glyph_and_missing_or_stale_evidence(self):
        f=self.fixture();out=f.build();summary=inspect_docx.run_inspection(out,f.root/'inspection','draft',True,True);report=read(summary['inspection'])
        bio_contract.check(report,f.e['title_candidate'])
        for kind in ('missing','stale','incomplete'):
            bad=deepcopy(report)
            if kind=='missing':bad.pop('biochemical_native_scripts_evidence')
            elif kind=='stale':bad['biochemical_native_scripts_evidence']['artifact_sha256']='0'*64
            else:bad['biochemical_native_scripts_evidence']['complete']=False
            with self.assertRaises(batch_tool.BatchError):bio_contract.check(bad,f.e['title_candidate'])
        d=Document(out);d.tables[0].cell(0,1).text='10⁻⁵';path=f.root/'tampered.docx';d.save(path)
        summary=inspect_docx.run_inspection(path,f.root/'tampered-inspection','draft',True,True);bad=read(summary['inspection'])
        self.assertTrue(bad['biochemical_native_scripts_evidence']['issues'])
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(bad,f.e['title_candidate'])

if __name__=='__main__':unittest.main()
