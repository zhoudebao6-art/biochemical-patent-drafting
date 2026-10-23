"""Executable domain regressions with isolated synthetic, non-experimental data."""
from pathlib import Path
from io import BytesIO
import tempfile, unittest, zipfile, copy, json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image
from lxml import etree as ET
from bio_common import sha, read, write
from bio_prepare import prepare, text
from bio_author import build
from bio_check import inspect
import bio_contract, inspect_docx, batch_tool

class BiochemicalBehaviorTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        self.source=self.root/'source.docx';d=Document()
        d.add_paragraph('技术项目优化说明书');d.add_paragraph('产品名称一种示例材料的制备方法')
        d.add_paragraph('技术概括');d.add_paragraph('本产品涉及示例材料技术领域，公开了一种示例材料的制备方法。保留样品身份。')
        d.add_paragraph('工作原理');d.add_paragraph('原理说明')
        d.add_paragraph('针对现有技术中样品身份不清的问题，本方案提供示例方法。')
        d.add_paragraph('为了解决上述技术问题，本发明提供了一种示例材料的制备方法，其特征在于，包括以下步骤：将样品A转移，得到试液；将所述试液检测。')
        d.add_paragraph('具体实施中，以下仅为软件测试夹具，不是实验结果。')
        d.add_paragraph('进一步地，在上述技术方案中，所述试液保存在样品瓶中。')
        d.add_paragraph('具体实施中，样品瓶的解释不作为新增权项。')
        d.add_paragraph('进一步地，在上述技术方案中，所述样品瓶具有标签。')
        d.add_paragraph('产品背景');d.add_paragraph('样品需要可追溯。')
        d.add_paragraph('技术效果');d.add_paragraph('本技术保持样品编号。')
        d.add_paragraph('实施方案');p=d.add_paragraph('保留(SO');p.add_run('4').font.subscript=True;p.add_run(')及cm');p.add_run('2').font.superscript=True;p.add_run('。')
        p=d.add_paragraph('原生数学：');math=OxmlElement('m:oMath');frac=OxmlElement('m:f')
        for tag,val in [('num','a'),('den','b')]:
            n=OxmlElement('m:'+tag);r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=val;r.append(t);n.append(r);frac.append(n)
        math.append(frac);p._p.append(math)
        table=d.add_table(rows=2,cols=2);table.cell(0,0).text='样品';table.cell(0,1).text='值';table.cell(1,0).text='A';table.cell(1,1).text='—'
        d.add_paragraph('结果分析：这是测试夹具。');d.add_paragraph('参考图示')
        for num in (1,2):
            b=BytesIO();Image.new('RGB',(100+num,80),'white').save(b,format='PNG');b.seek(0);d.add_picture(b,width=Cm(3))
            d.add_paragraph(('（主图）' if num==2 else '')+f'图{num}：示例图{num}。');d.add_paragraph(f'由图{num}可知，示例标签保留。')
        d.save(self.source);self.original=sha(self.source);self.e=prepare(self.source,self.root/'evidence')
    def plan(self):
        e=self.e;bysec=lambda h:[r for r in e['body'] if r['section']==h and r['text'].strip().rstrip('：:')!=h]
        abstract=bysec('技术概括')[0];claims=[];mapping=[]
        for i,c in enumerate(e['claim_candidates'],1):
            tx=c['text'];start=tx.index('其特征在于，')+len('其特征在于，') if i==1 else len('进一步地，在上述技术方案中，')
            parent=None if i==1 else (1 if i==2 else 2)
            obj='一种示例材料的制备方法'
            mapping.append({'number':i,'parent':parent,'object':obj,'source_body_indices':[c['body_index']],'dependency_reason':'试液在权1、样品瓶在权2独立引入'})
            if parent:
                mapping[-1]['dependency_basis']=[{'claim':parent,'body_index':e['claim_candidates'][parent-1]['body_index'],
                    'quote':'得到试液' if parent==1 else '所述试液保存在样品瓶中','current_quote':tx[start:]}]
            prefix=f'{i}、'+(f'根据权利要求{parent}所述的' if parent else '')+obj+'，其特征在于，'
            claims.append({'source_body_index':c['body_index'],'slice':[start,len(tx)],'prefix':prefix})
        desc=[{'text':e['title_candidate'],'role':'title'},{'text':'技术领域','role':'heading'},{'source_body_index':abstract['body_index'],'slice':[0,abstract['text'].index('。')+1]}]
        for h,source in [('背景技术','产品背景'),('发明内容','原理说明'),('具体实施方式','实施方案')]:
            if h=='具体实施方式':
                desc.extend({'source_body_index':r['body_index']} for r in bysec('技术效果'))
                desc.append({'text':'附图说明','role':'heading'})
                desc.extend({'text':f"图{f['number']}为本发明的{f['name']}。",'source_body_indices':[f['caption_body_index']]} for f in e['figures'])
            desc.append({'text':h,'role':'heading'})
            for r in bysec(source):
                if r['kind']=='tbl':desc.append({'text':'表1样品检测值','role':'table_caption','source_body_indices':[r['body_index']]})
                desc.append({'source_body_index':r['body_index']})
        for f in e['figures']:desc.extend({'source_body_index':i} for i in f['explanation_body_indices'])
        figs=[]
        for f in e['figures']:
            figs.extend({'source_body_index':i,'role':'image'} for i in f['image_body_indices']);figs.append({'text':'图'+str(f['number']),'role':'caption'})
        return {'schema_version':1,'source_sha256':e['source']['sha256'],'title':e['title_candidate'],'claim_map':mapping,'sections':{'abstract':[{'source_body_index':abstract['body_index']}],'abstract_figure':[{'source_body_index':e['figures'][1]['image_body_indices'][0],'role':'image'}],'claims':claims,'description':desc,'figures':figs}}
    def build(self):
        p=self.root/'plan.json';write(p,self.plan());out=self.root/'draft.docx';build(self.root/'evidence/evidence.json',p,out);return out
    def test_proposal_extraction_separates_explanations_and_finds_nonfirst_main(self):
        self.assertEqual(len(self.e['claim_candidates']),3)
        self.assertEqual([f['number'] for f in self.e['figures'] if f['main']],[2])
        self.assertEqual(self.original,sha(self.source))
    def test_rich_generation_preserves_math_scripts_tables_images_and_brackets(self):
        out=self.build();d=Document(out);src=Document(self.source)
        self.assertEqual(len(d.sections),5);self.assertEqual(len(d.inline_shapes),3)
        self.assertEqual([[c.text for c in r.cells] for r in d.tables[0].rows],[[c.text for c in r.cells] for r in src.tables[0].rows])
        self.assertIn('(SO4)',text(d.element.body));self.assertEqual(len(d.element.body.xpath('.//w:vertAlign')),2)
        self.assertEqual(ET.tostring(d.element.body.xpath('.//m:oMath')[0],method='c14n',exclusive=True),ET.tostring(src.element.body.xpath('.//m:oMath')[0],method='c14n',exclusive=True))
        self.assertEqual(self.original,sha(self.source))
        self.assertFalse(inspect(out,self.plan(),self.e)['errors'])
    def test_biochemical_pass_accepts_three_claims_and_required_brackets(self):
        out=self.build();summary=inspect_docx.run_inspection(out,self.root/'inspection','draft',True,True);report=read(summary['inspection'])
        result=bio_contract.check(report,self.e['title_candidate']);self.assertEqual(result['mode'],'biochemical_v1')
        self.assertFalse(any(x['code']=='CLAIM_SHUSU_MINIMUM' for x in report['issues']))
        self.assertEqual(len(report['media']),2)
        self.assertTrue(all(Path(m['extracted_path']).is_file() for m in report['media']))
        self.assertEqual(report['claims_detected'][2]['inheritance_paths'],[[3,2,1]])
    def test_stale_source_and_overwrite_are_rejected(self):
        out=self.build()
        with self.assertRaisesRegex(ValueError,'overwrite'):build(self.root/'evidence/evidence.json',self.root/'plan.json',out)
        self.source.write_bytes(self.source.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'identity'):build(self.root/'evidence/evidence.json',self.root/'plan.json',self.root/'other.docx')
    def test_distinct_tables_cannot_merge_when_plan_omits_blank_separator(self):
        plan=self.plan();blocks=plan['sections']['description'];idx=next(i for i,b in enumerate(blocks) if 'source_body_index' in b and self.e['body'][b['source_body_index']]['kind']=='tbl')
        blocks.insert(idx+1,copy.deepcopy(blocks[idx]));write(self.root/'plan.json',plan)
        out=self.root/'tables.docx';build(self.root/'evidence/evidence.json',self.root/'plan.json',out);d=Document(out)
        tables=d.element.body.xpath('./w:tbl');self.assertEqual(len(tables),2)
        self.assertEqual(ET.QName(tables[0].getnext()).localname,'p')
        self.assertEqual(ET.QName(tables[0].getnext().getnext()).localname,'tbl')
    def test_wrong_parent_and_unresolved_feature_phrase_are_not_passed(self):
        plan=self.plan();plan['claim_map'][1]['parent']=3;write(self.root/'badplan.json',plan)
        with self.assertRaisesRegex(ValueError,'parent'):build(self.root/'evidence/evidence.json',self.root/'badplan.json',self.root/'bad.docx')
        out=self.build();summary=inspect_docx.run_inspection(out,self.root/'inspection','draft',True,True);report=read(summary['inspection'])
        bad=copy.deepcopy(report);bad['biochemical_evidence']['wording_hits']=[{'word':'其特征在于'}]
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(bad,self.e['title_candidate'])
if __name__=='__main__':unittest.main()
