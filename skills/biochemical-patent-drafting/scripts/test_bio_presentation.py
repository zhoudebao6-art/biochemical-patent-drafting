"""Observable heading/caption regressions, including source and style inheritance."""
import copy,unittest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
import test_biochemical as fixtures
from bio_common import read,write
from bio_author import build
from bio_layout import normalize_plan,preview
from bio_presentation import BENEFIT_INTRO,CAPTION,evidence as presentation,visible,bold_reader
import inspect_docx,bio_contract,batch_tool

class PresentationTests(unittest.TestCase):
    def fixture(self):
        f=fixtures.BiochemicalBehaviorTests();f.setUp();self.addCleanup(f.doCleanups);return f
    def test_bad_extra_headings_and_intro_become_one_plain_introduction_with_facts_preserved(self):
        f=self.fixture();plan=f.plan();blocks=plan['sections']['description'];idx=next(i for i,b in enumerate(blocks) if b.get('source_body_index')==next(r['body_index'] for r in f.e['body'] if r['text']=='本技术保持样品编号。'))
        support=blocks[idx]['source_body_index']
        blocks[idx:idx]=[{'text':'有益效果','role':'heading'},{'text':'本发明带来的技术效果如下：','role':'heading','source_body_indices':[support]},{'text':'本发明的有益效果如下：','role':'heading','source_body_indices':[support]},{'text':BENEFIT_INTRO,'role':'heading','source_body_indices':[support]}]
        write(f.root/'plan.json',plan);out=f.root/'draft.docx';build(f.root/'evidence/evidence.json',f.root/'plan.json',out)
        d=Document(out);texts=[p.text for p in d.paragraphs]
        self.assertEqual(texts.count(BENEFIT_INTRO),1);self.assertNotIn('有益效果',texts);self.assertNotIn('本发明的有益效果如下：',texts);self.assertIn('本发明保持样品编号。',texts)
        p=next(p for p in d.paragraphs if p.text==BENEFIT_INTRO);self.assertTrue(all(r.bold is False for r in p.runs))
    def test_inherited_bold_is_explicitly_cleared_for_preparation_label_and_caption(self):
        f=self.fixture();plan=f.plan();idx=next(i for i,b in enumerate(plan['sections']['description']) if b.get('role')=='table_caption');support=plan['sections']['description'][idx]['source_body_indices']
        plan['sections']['description'].insert(idx,{'text':'制备工艺：','role':'heading','source_body_indices':support})
        write(f.root/'plan.json',plan);out=f.root/'draft.docx';build(f.root/'evidence/evidence.json',f.root/'plan.json',out)
        d=Document(out)
        for p in d.paragraphs:
            if p.text=='制备工艺：' or p.text.startswith('表1'):
                self.assertEqual(p.style.name,'Normal');self.assertTrue(all(r.bold is False for r in p.runs))
        heading=next(p for p in d.paragraphs if p.text=='具体实施方式');self.assertTrue(all(r.bold for r in heading.runs))
    def test_every_table_gets_numbered_nonbold_caption_from_source_shared_heading(self):
        e={'body':[{'body_index':0,'text':'各实施例与对比例的性能测试结果','kind':'p'},{'body_index':1,'text':'表值甲','kind':'tbl'},{'body_index':2,'text':'表值乙','kind':'tbl'}]}
        p={'sections':{'description':[{'source_body_index':i} for i in range(3)]}}
        norm,_=normalize_plan(p,e);blocks=norm['sections']['description']
        self.assertEqual([preview(b,e) for b in blocks if b.get('role')=='table_caption'],['表1各实施例与对比例的性能测试结果1','表2各实施例与对比例的性能测试结果2'])
        self.assertEqual(blocks[0]['text'],'测试结果如表1至表2所示。')
        self.assertEqual([b['source_body_index'] for b in blocks if 'source_body_index' in b],[1,2])
    def test_existing_source_caption_is_not_duplicated_and_references_are_not_silently_renumbered(self):
        e={'body':[{'body_index':0,'text':'表1原始指标说明','kind':'p'},{'body_index':1,'text':'原数据','kind':'tbl'}]}
        p={'sections':{'description':[{'source_body_index':0},{'source_body_index':1}]}}
        norm,_=normalize_plan(p,e);self.assertEqual(len(norm['sections']['description']),2)
        p['sections']['description'][0]={'text':'表8原始指标说明','source_body_indices':[0],'role':'table_caption'}
        with self.assertRaisesRegex(ValueError,'renumbering'):normalize_plan(p,e)
    def test_performance_caption_group_does_not_leak_into_later_unrelated_tables(self):
        e={'body':[{'body_index':0,'text':'各实施例与对比例的性能测试结果','kind':'p'},{'body_index':1,'text':'性能数据','kind':'tbl'},{'body_index':2,'text':'结果分析：','kind':'p'},{'body_index':3,'text':'另有配方表','kind':'tbl'}]}
        p={'sections':{'description':[{'source_body_index':i} for i in range(4)]}}
        with self.assertRaisesRegex(ValueError,'source-bound descriptive caption'):normalize_plan(p,e)
    def test_bare_number_and_punctuation_are_not_descriptive_table_captions(self):
        for value in ('表1','表12','表123','表1：','表12:','表1： '):
            with self.subTest(value=value):self.assertIsNone(CAPTION.fullmatch(value))
        self.assertEqual(CAPTION.fullmatch('表12：各样品检测结果').groups(),('12','各样品检测结果'))
    def test_omitted_source_caption_cannot_silently_renumber_source_references(self):
        e={'body':[{'body_index':0,'text':'表8原始指标说明','kind':'p'},{'body_index':1,'text':'原数据','kind':'tbl'}]}
        p={'sections':{'description':[{'source_body_index':1}]}}
        with self.assertRaisesRegex(ValueError,'renumbering'):normalize_plan(p,e)
    def test_unavailable_caption_support_and_invented_extra_heading_are_rejected(self):
        f=self.fixture();p=f.plan();p['sections']['description']=[b for b in p['sections']['description'] if b.get('role')!='table_caption']
        with self.assertRaisesRegex(ValueError,'source-bound descriptive caption'):normalize_plan(p,f.e)
        p=f.plan();p['sections']['description'].append({'text':'额外总结标题','role':'heading'});write(f.root/'plan.json',p)
        with self.assertRaisesRegex(ValueError,'source locators'):build(f.root/'evidence/evidence.json',f.root/'plan.json',f.root/'bad.docx')
    def test_real_inspector_rejects_missing_second_caption_and_inherited_label_bold(self):
        f=self.fixture();out=f.build();d=Document(out);first=d.tables[0]._tbl
        second=copy.deepcopy(first);first.addnext(second)
        style=d.styles.add_style('UnexpectedBoldLabel',WD_STYLE_TYPE.PARAGRAPH);style.font.bold=True
        label=next(p for p in d.paragraphs if p.text==BENEFIT_INTRO);label.style=style
        for run in label.runs:
            run.font.bold=None
            cs=run._r.rPr.find(qn('w:bCs'))
            if cs is not None:run._r.rPr.remove(cs)
        tampered=f.root/'tampered.docx';d.save(tampered)
        summary=inspect_docx.run_inspection(tampered,f.root/'inspection','draft',True,True);report=read(summary['inspection'])
        codes={i['code'] for i in report['biochemical_presentation_evidence']['issues']}
        self.assertIn('descriptive_table_caption_missing',codes);self.assertIn('non_main_text_bold',codes)
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(report,f.e['title_candidate'])
    def test_header_rows_stay_with_first_data_row(self):
        f=self.fixture();out=f.build();d=Document(out)
        header=d.tables[0].rows[0];data=d.tables[0].rows[1]
        self.assertTrue(all(p.paragraph_format.keep_with_next for c in header.cells for p in c.paragraphs))
        self.assertTrue(all(p.paragraph_format.keep_with_next is False for c in data.cells for p in c.paragraphs))
        caption=next(p for p in d.paragraphs if p.text.startswith('表1'));self.assertTrue(caption.paragraph_format.keep_with_next)
        label=next(p for p in d.paragraphs if p.text=='结果分析：');self.assertTrue(label.paragraph_format.keep_with_next);self.assertTrue(all(r.bold is False for r in label.runs))
    def test_final_figure_description_ends_in_full_stop_without_touching_analysis_or_figure_numbers(self):
        f=self.fixture();out=f.build();d=Document(out)
        start=next(i for i,p in enumerate(d.paragraphs) if p.text=='附图说明')
        descriptions=d.paragraphs[start+1:start+3]
        self.assertEqual([p.text for p in descriptions],['图1为本发明的示例图1；','图2为本发明的示例图2。'])
        self.assertIn('由图1可知，示例标签保留。',[p.text for p in d.paragraphs])
        self.assertIn('图1',[p.text for p in d.paragraphs])
        self.assertEqual(len(read(out.with_suffix('.trace.json'))['figure_description_endings']),1)
        descriptions[-1].runs[-1].text=descriptions[-1].runs[-1].text.rstrip('。')+'；'
        tampered=f.root/'bad-figure-ending.docx';d.save(tampered)
        summary=inspect_docx.run_inspection(tampered,f.root/'end-inspection','draft',True,True);report=read(summary['inspection'])
        self.assertIn('figure_description_terminal_not_full_stop',[i['code'] for i in report['biochemical_presentation_evidence']['issues']])
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(report,f.e['title_candidate'])
    def test_terminal_edit_preserves_rich_scripts_decimal_points_and_internal_sentences(self):
        from bio_presentation import normalize_figure_description_endings
        d=Document();d.add_paragraph('附图说明');p=d.add_paragraph('图1为本发明在0.5条件下的图。内部说明。')
        q=d.add_paragraph('图2为本发明的cm');q.add_run('2').font.superscript=True
        d.add_paragraph('图3为本发明的图；');d.add_paragraph('具体实施方式');d.add_paragraph('图1对应的数据为0.5。')
        normalize_figure_description_endings(d.element)
        self.assertEqual(p.text,'图1为本发明在0.5条件下的图。内部说明；')
        self.assertEqual(q.text,'图2为本发明的cm2；');self.assertTrue(q.runs[1].font.superscript)
        self.assertFalse(q.runs[-1].font.superscript);self.assertEqual(d.paragraphs[-1].text,'图1对应的数据为0.5。')
        self.assertEqual(d.paragraphs[3].text,'图3为本发明的图。')
        self.assertFalse(normalize_figure_description_endings(d.element))
    def test_single_or_empty_figure_description_list(self):
        from bio_presentation import normalize_figure_description_endings,figure_description_entries
        d=Document();d.add_paragraph('附图说明')
        self.assertEqual(normalize_figure_description_endings(d.element),[])
        p=d.add_paragraph('图7为本发明的唯一示意图；')
        self.assertEqual(figure_description_entries(d.element)[0][2],'。')
        self.assertEqual(len(normalize_figure_description_endings(d.element)),1)
        self.assertEqual(p.text,'图7为本发明的唯一示意图。')
        self.assertFalse(normalize_figure_description_endings(d.element))

    def test_each_list_uses_actual_final_entry_across_blanks_and_section_boundaries(self):
        from bio_presentation import normalize_figure_description_endings
        d=Document();d.add_paragraph('附图说明：')
        first=d.add_paragraph('图3为本发明的前图。');d.add_paragraph('')
        final=d.add_paragraph('图8为本发明的后图;');d.add_paragraph('')
        d.add_paragraph('具体实施方式');body=d.add_paragraph('图9为本发明的结果分析；')
        d.add_paragraph('附图说明');only=d.add_paragraph('图1为本发明的另一列表；')
        normalize_figure_description_endings(d.element)
        self.assertEqual([p.text[-1] for p in (first,final,body,only)],['；','。','；','。'])

    def test_nonfinal_full_stop_is_still_rejected(self):
        f=self.fixture();out=f.build();d=Document(out)
        p=next(p for p in d.paragraphs if p.text=='图1为本发明的示例图1；')
        p.runs[-1].text=p.runs[-1].text.rstrip('；')+'。'
        errors=presentation(d.element,d.styles.element,'fixture')['issues']
        self.assertIn('figure_description_terminal_not_semicolon',[i['code'] for i in errors])
    def examples_fixture(self):
        from bio_prepare import prepare
        from bio_common import sha
        f=self.fixture();d=Document(f.source)
        anchor=next(p for p in d.paragraphs if p.text.startswith('结果分析'))
        anchor.insert_paragraph_before('实施例1：本实施例为一种示例材料的制备方法的优选实施例。')
        p=anchor.insert_paragraph_before('配方：按来源保存Na');p.add_run('2').font.subscript=True;p.add_run('CO');p.add_run('3').font.subscript=True;p.add_run('。')
        anchor.insert_paragraph_before('制备工艺：')
        anchor.insert_paragraph_before('S1、按原文记录样品。')
        anchor.insert_paragraph_before('实施例2：与实施例1的不同在于，保存另一原有样品编号。')
        anchor.insert_paragraph_before('对比例1：保持来源的对照条件。')
        d.save(f.source);f.original=sha(f.source);f.e=prepare(f.source,f.root/'examples-evidence')
        return f

    def test_example_source_sentence_cannot_be_omitted_truncated_duplicated_or_split(self):
        f=self.examples_fixture();plan=f.plan()
        from bio_presentation import example_entries
        cases=example_entries([{'text':r['text'],'location':r['body_index'],'kind':r['kind']} for r in f.e['body']],True)
        idx=cases[0]['locations'][0];full=f.e['body'][idx]['text']
        normalize_plan(plan,f.e)
        for kind in ('omit','bare','truncated','duplicate','split','later_difference'):
            p=copy.deepcopy(plan);blocks=p['sections']['description'];pos=next(i for i,b in enumerate(blocks) if b.get('source_body_index')==idx)
            if kind=='omit':blocks.pop(pos)
            elif kind=='bare':blocks[pos]['slice']=[0,len('实施例1')]
            elif kind=='truncated':blocks[pos]['replace']=[['的优选实施例。','。']]
            elif kind=='duplicate':blocks.insert(pos,copy.deepcopy(blocks[pos]))
            elif kind=='split':blocks[pos:pos+1]=[{'source_body_index':idx,'slice':[0,len('实施例1：')]},{'source_body_index':idx,'slice':[len('实施例1：'),len(full)]}]
            else:
                target=next(b for b in blocks if b.get('source_body_index')==cases[1]['locations'][0]);target['replace']=[['另一原有样品编号','样品']]
            with self.subTest(kind=kind),self.assertRaisesRegex(ValueError,'Source example introduction'):
                normalize_plan(p,f.e)

    def test_example_identification_excludes_results_references_and_preserves_comparative_identity(self):
        from bio_presentation import example_entries,compare_examples
        raw=['实施方案','实施例1：本实施例为本方案的优选实施例。','实施例1至实施例5的结果如下。','对比例1：保留原对照条件。','参考图示','实施例2：不属于实施方案。']
        src=example_entries([{'text':t,'location':i} for i,t in enumerate(raw)],True)
        self.assertEqual([r['key'] for r in src],['实施例1','对比例1'])
        self.assertIn('本发明',src[0]['introduction']);self.assertNotIn('优选',src[1]['introduction'])
        self.assertFalse(compare_examples(src,src))
        self.assertTrue(compare_examples(src,[]))

    def test_source_sentence_and_native_scripts_survive_real_build_with_role_geometry(self):
        from docx.shared import Pt
        f=self.examples_fixture();write(f.root/'plan.json',f.plan());out=f.root/'examples.docx'
        build(f.root/'examples-evidence/evidence.json',f.root/'plan.json',out);d=Document(out)
        intro=next(p for p in d.paragraphs if p.text.startswith('实施例1：'))
        recipe=next(p for p in d.paragraphs if p.text.startswith('配方：'))
        for p in (intro,recipe):
            self.assertEqual(p.paragraph_format.line_spacing,1.3)
            self.assertEqual(p.paragraph_format.space_before.pt,5)
            self.assertTrue(all(r.bold is False and r.font.size.pt==12 for r in p.runs))
        self.assertEqual([r.text for r in recipe.runs if r.font.subscript],['2','3'])
        claim=next(p for p in d.paragraphs if p.text.startswith('1、'));self.assertEqual(claim.paragraph_format.line_spacing.pt,24)
        report=read(inspect_docx.run_inspection(out,f.root/'examples-scan','draft',False,True)['inspection'])
        source=read(inspect_docx.run_inspection(f.source,f.root/'source-scan','proposal',False,True)['inspection'])
        bio_contract.check(report,f.e['title_candidate'],source_reports=[source],require_example_sources=True)
        for kind in ('delete_all','truncate','line_spacing'):
            changed=Document(out)
            if kind=='delete_all':
                for p in list(changed.paragraphs):
                    if p.text.startswith(('实施例','对比例')):p._p.getparent().remove(p._p)
            elif kind=='truncate':next(p for p in changed.paragraphs if p.text.startswith('实施例1：')).runs[0].text='实施例1：本实施例为一种示例材料的制备方法。'
            else:next(p for p in changed.paragraphs if p.text.startswith('实施例1：')).paragraph_format.line_spacing=Pt(24)
            path=f.root/(kind+'.docx');changed.save(path)
            bad=read(inspect_docx.run_inspection(path,f.root/(kind+'-scan'),'draft',False,True)['inspection'])
            with self.subTest(kind=kind),self.assertRaises(batch_tool.BatchError):
                bio_contract.check(bad,f.e['title_candidate'],source_reports=[source],require_example_sources=True)
        with self.assertRaisesRegex(batch_tool.BatchError,'source reports'):
            bio_contract.check(report,f.e['title_candidate'],require_example_sources=True)

    def test_description_geometry_merges_style_attributes_and_checks_effective_side_indents(self):
        from docx.oxml import OxmlElement
        f=self.fixture();out=f.build();d=Document(out)
        p=next(p for p in d.paragraphs if p.text.startswith('保留(SO'))
        style=d.styles.add_style('InheritedBody',WD_STYLE_TYPE.PARAGRAPH);style.base_style=d.styles['Normal']
        props=style.element.get_or_add_pPr()
        for name,attrs in [('spacing',{'before':'100','after':'0','line':'312','lineRule':'auto'}),('ind',{'firstLineChars':'200'})]:
            item=OxmlElement('w:'+name)
            for key,value in attrs.items():item.set(qn('w:'+key),value)
            props.append(item)
            direct=p._p.pPr.find(qn('w:'+name));p._p.pPr.remove(direct)
        p.style=style;direct=OxmlElement('w:ind');direct.set(qn('w:right'),'0');p._p.pPr.append(direct)
        self.assertFalse(any(i['code'].startswith('description_body') for i in presentation(d.element,d.styles.element,'a'*64)['issues']))
        for attr,value in [('start','200'),('endChars','100'),('firstLineChars','0')]:
            direct.set(qn('w:'+attr),value)
            with self.subTest(attr=attr):
                self.assertTrue(any(i['code']=='description_body_paragraph_geometry' for i in presentation(d.element,d.styles.element,'a'*64)['issues']))
            del direct.attrib[qn('w:'+attr)]

    def test_example_comparison_allows_only_explicit_native_script_mapping(self):
        from bio_presentation import example_entries,compare_examples
        def rows(heading,text):return [{'text':heading,'location':0},{'text':text,'location':1}]
        source=example_entries(rows('实施方案','实施例1：本实施例为LiFePO₄及10⁻⁵条件的优选实施例。'),True)
        final=example_entries(rows('具体实施方式','实施例1：本实施例为LiFePO4及10-5条件的优选实施例。'))
        self.assertFalse(compare_examples(source,final))
        self.assertIn('₄',source[0]['raw_text'])
        altered=example_entries(rows('具体实施方式','实施例1：本实施例为LiFePO及105条件的优选实施例。'))
        self.assertTrue(compare_examples(source,altered))
        self.assertTrue(compare_examples(example_entries(rows('实施方案','实施例1：保留ℬ符号。'),True),example_entries(rows('具体实施方式','实施例1：保留B符号。'))))

    def test_final_gate_reuses_verified_source_objects_without_reparsing(self):
        from unittest.mock import patch
        import output_contract
        f=self.examples_fixture();write(f.root/'plan.json',f.plan());out=f.root/'examples.docx'
        build(f.root/'examples-evidence/evidence.json',f.root/'plan.json',out)
        report=read(inspect_docx.run_inspection(out,f.root/'scan','draft',False,True)['inspection'])
        source=read(inspect_docx.run_inspection(f.source,f.root/'src-scan','proposal',False,True)['inspection'])
        verified={'handoff':{'input_inspection':{'sealed':True}},'input_report':report,'source_reports':[source]}
        m={'rule_version':8,'policy':{'biochemical_contract':'biochemical_v1','output_contract':output_contract.MODE}}
        with patch('prepare_handoff.verify_prepared',return_value=verified),patch.object(inspect_docx,'inspect',side_effect=AssertionError('No repeat parse')):
            output_contract.verify(m,{'title':f.e['title_candidate']},{'input_artifact_sha256':report['sha256']},{})
            report['biochemical_presentation_evidence']['examples']=[]
            with self.assertRaises(batch_tool.BatchError):
                output_contract.verify(m,{'title':f.e['title_candidate']},{'input_artifact_sha256':report['sha256']},{})



class ProseFlowTests(unittest.TestCase):
    def setUp(self):
        from bio_presentation import normalize_prose_flow
        self.normalize=normalize_prose_flow
    def fixture(self):
        f=fixtures.BiochemicalBehaviorTests();f.setUp();self.addCleanup(f.doCleanups);return f
    def paragraph(self):
        d=Document();p=d.add_paragraph('测试句中的值为1.25mAh·g');p.add_run('-1').font.superscript=True;p.add_run('，')
        return d,p
    def test_internal_line_page_and_carriage_breaks_preserve_rich_runs_and_inline_math(self):
        from docx.oxml import OxmlElement
        from lxml import etree
        for tag,typ in [('br','textWrapping'),('br','page'),('cr',None)]:
            with self.subTest(tag=tag,typ=typ):
                d,p=self.paragraph();br=OxmlElement('w:'+tag)
                if typ:br.set(qn('w:type'),typ)
                p.add_run()._r.append(br);p.add_run('另一值为2.50mAh·cm');p.add_run('-2').font.superscript=True
                math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x';r.append(t);math.append(r);p._p.append(math)
                before=visible(p._p);math_xml=etree.tostring(math)
                self.assertEqual(len(self.normalize(p._p,'body','description')),1)
                self.assertEqual(visible(p._p),before);self.assertEqual(etree.tostring(math),math_xml)
                self.assertEqual([r.text for r in p.runs if r.font.superscript],['-1','-2'])
                self.assertFalse(p._p.xpath('.//w:br | .//w:cr'));self.assertEqual(self.normalize(p._p,'body','description'),[])
    def test_ambiguous_break_is_rejected_without_partial_changes(self):
        from docx.enum.text import WD_BREAK
        for tail in ('S2、开始下一步骤。','另一个没有前导逗号的语句。'):
            d,p=self.paragraph();p.add_run().add_break();p.add_run(tail);p.add_run().add_break(WD_BREAK.COLUMN);p.add_run('后文。')
            before=p._p.xml
            with self.assertRaisesRegex(ValueError,'Ambiguous manual prose break'):self.normalize(p._p,'body','description')
            self.assertEqual(p._p.xml,before)
    def test_other_scopes_and_cached_pagination_are_preserved(self):
        from docx.oxml import OxmlElement
        d,p=self.paragraph();p.add_run().add_break();p.add_run('后文。')
        for role,slot in [('body','claims'),('table','description'),('caption','description'),('image','figures')]:
            before=p._p.xml;self.assertFalse(self.normalize(p._p,role,slot));self.assertEqual(p._p.xml,before)
        d,p=self.paragraph();p.add_run()._r.append(OxmlElement('w:lastRenderedPageBreak'));p.add_run('后文。')
        before=p._p.xml;self.assertFalse(self.normalize(p._p,'body','description'));self.assertEqual(p._p.xml,before)
    def test_literal_newline_is_not_silently_removed(self):
        d,p=self.paragraph();p.runs[-1]._r.find(qn('w:t')).text='，\n后文。';before=p._p.xml
        with self.assertRaisesRegex(ValueError,'Ambiguous manual prose break'):self.normalize(p._p,'body','description')
        self.assertEqual(p._p.xml,before)
    def test_body_pagination_overrides_style_but_heading_stays_with_next(self):
        from bio_author import format_p
        d,p=self.paragraph();d.styles['Normal'].paragraph_format.keep_with_next=True
        format_p(p._p,'body','description')
        for flag in ('keepNext','keepLines','pageBreakBefore','widowControl'):self.assertEqual(p._p.find('w:pPr/w:'+flag,p._p.nsmap).get(qn('w:val')),'0')
        heading=d.add_paragraph('技术领域');format_p(heading._p,'heading','description')
        self.assertTrue(heading.paragraph_format.keep_with_next)
    def test_inspection_flags_breaks_split_paragraphs_and_effective_pagination(self):
        from bio_author import format_p
        d=Document();d.add_paragraph('一种格式测试材料');d.add_paragraph('技术领域')
        a=d.add_paragraph('前半句，');b=d.add_paragraph('后半句。')
        for p in (a,b):format_p(p._p,'body','description')
        b.add_run().add_break();b.add_run('仍有内容。');a.paragraph_format.keep_with_next=True
        codes={i['code'] for i in presentation(d.element,d.styles.element,'fixture')['issues']}
        self.assertTrue({'description_manual_break','description_prose_forced_pagination','description_prose_split_after_comma'}<=codes)
        b.insert_paragraph_before('')
        split=next(x for x in presentation(d.element,d.styles.element,'fixture')['issues'] if x['code']=='description_prose_split_after_comma')
        self.assertEqual(split['next_body_index']-split['body_index'],2)
        a.paragraph_format.keep_with_next=None;d.styles['Normal'].paragraph_format.keep_with_next=True
        self.assertIn('description_prose_forced_pagination',{i['code'] for i in presentation(d.element,d.styles.element,'fixture')['issues']})
    def test_numbered_steps_and_section_boundaries_are_not_merged(self):
        from bio_author import format_p
        from docx.enum.section import WD_SECTION_START
        d=Document();d.add_paragraph('一种格式测试材料');d.add_paragraph('技术领域')
        for t in ('S1、第一独立步骤，','S2、第二独立步骤。'):
            p=d.add_paragraph(t);format_p(p._p,'body','description')
        d.add_section(WD_SECTION_START.NEW_PAGE);d.add_paragraph('附图占位')
        ev=presentation(d.element,d.styles.element,'fixture')
        self.assertNotIn('description_prose_split_after_comma',{x['code'] for x in ev['issues']})
        self.assertEqual(len(d.sections),2)
    def test_real_build_normalizes_copied_source_break_and_keeps_source_read_only(self):
        from bio_prepare import prepare
        from bio_common import sha
        f=self.fixture();d=Document(f.source);p=next(p for p in d.paragraphs if p.text.startswith('结果分析'))
        anchor=p;p=d.add_paragraph();anchor._p.addnext(p._p);p.add_run('测试值为1.25mAh·g');p.add_run('-1').font.superscript=True;p.add_run('，');p.add_run().add_break();p.add_run('另一值为2.50mAh·cm');p.add_run('-2').font.superscript=True;p.add_run('。')
        d.save(f.source);source_sha=sha(f.source);f.e=prepare(f.source,f.root/'flow-evidence');write(f.root/'plan.json',f.plan())
        out=f.root/'flow.docx';build(f.root/'flow-evidence/evidence.json',f.root/'plan.json',out)
        self.assertEqual(sha(f.source),source_sha)
        current=next(p for p in Document(out).paragraphs if '测试值为' in p.text)
        self.assertNotIn('\n',current.text);self.assertIn('，另一值为',current.text)
        self.assertEqual([r.text for r in current.runs if r.font.superscript],['-1','-2'])
        self.assertTrue(any(b.get('prose_flow_normalization') for b in read(out.with_suffix('.trace.json'))['blocks']))
    def test_final_contract_rejects_old_presentation_evidence(self):
        f=self.fixture();out=f.build();info=inspect_docx.run_inspection(out,f.root/'flow-inspection','draft',True,True);report=read(info['inspection'])
        bio_contract.check(report,f.e['title_candidate'])
        report['biochemical_presentation_evidence'].pop('prose_flow_contract_version')
        with self.assertRaisesRegex(batch_tool.BatchError,'continuous prose flow evidence missing'):bio_contract.check(report,f.e['title_candidate'])

if __name__=='__main__':unittest.main()
