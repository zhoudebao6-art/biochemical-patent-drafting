"""No-space output, source preservation, native objects and real pass-gate tests."""
from pathlib import Path
import unittest,copy,zipfile
from docx import Document
from docx.oxml import OxmlElement
from lxml import etree as ET
import test_biochemical as fixtures
from bio_prepare import prepare,text
from bio_common import sha,read
from bio_author import build
from bio_check import inspect
import bio_contract,inspect_docx,batch_tool
from bio_spacing import scan_docx,strip,W,M

class NoSpaceBehaviorTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.BiochemicalBehaviorTests('test_proposal_extraction_separates_explanations_and_finds_nonfirst_main')
        self.f.setUp();self.addCleanup(self.f.doCleanups)
    def replace_source(self,doc):
        doc.save(self.f.source);self.f.original=sha(self.f.source);self.f.e=prepare(self.f.source,self.f.root/'new-source-evidence')
        # The fixture's plan/build methods resolve this canonical file location.
        from bio_common import write
        write(self.f.root/'evidence/evidence.json',self.f.e)
    def test_unicode_spaces_removed_across_runs_tables_headers_math_without_source_edits(self):
        d=Document(self.f.source);p=next(p for p in d.paragraphs if p.text.startswith('保留('))
        p.add_run('浓度 ');p.add_run('12.5\u00a0mg/L\u3000记录\u200b。')
        d.tables[0].cell(0,0).text='样\u202f品'
        math=d.element.body.xpath('.//m:oMath')[0];math.findall('.//{'+M+'}t')[0].text=' a '
        self.replace_source(d);before=self.f.source.read_bytes();out=self.f.build();doc=Document(out)
        self.assertEqual(scan_docx(out)['hits'],[])
        self.assertIn('浓度12.5mg/L记录。',text(doc.element.body));self.assertEqual(doc.tables[0].cell(0,0).text,'样品')
        self.assertEqual([n.text for n in doc.element.body.xpath('.//m:t')],['a','b'])
        self.assertEqual(len(doc.element.body.xpath('.//w:vertAlign')),2)
        self.assertEqual(self.f.source.read_bytes(),before)
        trace=read(out.with_suffix('.trace.json'));self.assertGreater(trace['spacing']['removed_characters'],0)
        self.assertEqual(len(doc.inline_shapes),3)
    def test_field_instruction_spaces_and_paragraph_structure_are_preserved(self):
        out=self.f.build();d=Document(out)
        self.assertEqual(len(d.sections),5)
        codes=[]
        for s in d.sections:codes.extend(n.text for n in s.footer._element.xpath('.//w:instrText'))
        self.assertTrue(codes);self.assertTrue(all(c==' PAGE ' for c in codes))
        self.assertIn('说明书',[p.text for s in d.sections for p in s.header.paragraphs])
        self.assertFalse(scan_docx(out)['hits'])
    def test_meaningful_word_number_and_variable_separation_cannot_be_blindly_joined(self):
        for value in ['sodium chloride','1 2']:
            with self.subTest(value=value):
                plan=self.f.plan();plan['sections']['description'].append({'text':value,'source_body_indices':[17]})
                from bio_common import write
                pp=self.f.root/(value.replace(' ','_')+'.json');write(pp,plan);out=self.f.root/(value.replace(' ','_')+'.docx')
                with self.assertRaisesRegex(ValueError,'source-bound rewording'):build(self.f.root/'evidence/evidence.json',pp,out)
                self.assertFalse(out.exists())
    def test_new_output_spaces_in_header_or_table_are_rejected_before_delivery(self):
        out=self.f.build();d=Document(out);d.sections[0].header.paragraphs[0].add_run('标\u00a0题');d.tables[0].cell(1,1).text='1 2'
        dirty=self.f.root/'dirty.docx';d.save(dirty)
        self.assertTrue(inspect(dirty,self.f.plan(),self.f.e)['errors'])
        summary=inspect_docx.run_inspection(dirty,self.f.root/'dirty-cache','draft',True,True);report=read(summary['inspection'])
        self.assertTrue(report['biochemical_spacing_evidence']['hits'])
        with self.assertRaisesRegex(batch_tool.BatchError,'no-space'):bio_contract.check(report,self.f.e['title_candidate'])
    def test_missing_or_stale_spacing_scan_cannot_pass(self):
        out=self.f.build();summary=inspect_docx.run_inspection(out,self.f.root/'cache','draft',True,True);report=read(summary['inspection'])
        bio_contract.check(report,self.f.e['title_candidate'])
        for mutation in ['missing','stale']:
            bad=copy.deepcopy(report)
            if mutation=='missing':bad.pop('biochemical_spacing_evidence')
            else:bad['biochemical_spacing_evidence']['artifact_sha256']='0'*64
            with self.assertRaises(batch_tool.BatchError):bio_contract.check(bad,self.f.e['title_candidate'])
    def test_math_token_separation_is_not_misread_as_padding(self):
        d=Document(self.f.source);math=d.element.body.xpath('.//m:oMath')[0]
        math.findall('.//{'+M+'}t')[0].text='a b';self.replace_source(d)
        before=self.f.source.read_bytes()
        with self.assertRaisesRegex(ValueError,'source-bound rewording'):self.f.build()
        self.assertEqual(self.f.source.read_bytes(),before)
    def test_editable_chart_numbering_and_vml_labels_but_not_formula_codes_are_cleaned(self):
        from bio_spacing import normalize,A,C,V
        out=self.f.build()
        with zipfile.ZipFile(out) as z:parts={n:z.read(n) for n in z.namelist()}
        xml=f'<c:chartSpace xmlns:c="{C}" xmlns:a="{A}"><a:p><a:r><a:t>图\u30001</a:t></a:r></a:p><c:v>样\u00a0品</c:v><c:f>Sheet 1!A1</c:f></c:chartSpace>'
        parts['word/charts/spacing-test.xml']=xml.encode('utf-8')
        parts['word/drawings/spacing-test.vml']=f'<v:shape xmlns:v="{V}"><v:textpath string="图\u202f1"/></v:shape>'.encode('utf-8')
        dirty=self.f.root/'drawing-labels.docx'
        with zipfile.ZipFile(dirty,'w') as z:
            for n,b in parts.items():z.writestr(n,b)
        self.assertTrue(scan_docx(dirty)['hits']);normalize(dirty);self.assertFalse(scan_docx(dirty)['hits'])
        with zipfile.ZipFile(dirty) as z:
            root=ET.fromstring(z.read('word/charts/spacing-test.xml'))
            self.assertEqual(root.find('.//{'+C+'}f').text,'Sheet 1!A1')
            self.assertEqual(root.find('.//{'+C+'}v').text,'样品')
if __name__=='__main__':unittest.main()
