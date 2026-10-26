"""Build a Word draft from an explicit source-bound plan, without inventing facts."""
from pathlib import Path, PurePosixPath
from copy import deepcopy
import argparse, re, uuid, tempfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from lxml import etree as ET
from bio_common import SKILL, read, write, sha
from bio_prepare import text
from bio_layout import normalize_plan, table_format, summary_format, SUMMARY
from bio_presentation import paragraph_role

SLOTS=['abstract','abstract_figure','claims','description','figures']
REPLACE=[('本产品','本发明'),('本方案','本发明'),('本技术','本发明')]
def format_p(p,role='body',slot=None):
    pr=p.get_or_add_pPr()
    if role in ('figure_analysis','results_analysis_heading','technical_summary'):
        pr.get_or_add_pStyle().val={'figure_analysis':'BioFigureAnalysis','results_analysis_heading':'BioResultsHeading','technical_summary':'BioTechnicalSummary'}[role]
    else:pr.get_or_add_pStyle().val='Normal'
    for tag in ['spacing','ind','jc','keepNext','keepLines','pageBreakBefore','widowControl','outlineLvl']:
        el=pr.find(qn('w:'+tag))
        if el is not None: pr.remove(el)
    for tag in ['keepLines','pageBreakBefore','widowControl']:
        e=OxmlElement('w:'+tag); e.set(qn('w:val'),'0'); pr.append(e)
    natural=slot=='description' or role in ('image','heading','title','caption')
    e=OxmlElement('w:spacing'); e.set(qn('w:before'),'100' if role not in ('image','caption') else '0'); e.set(qn('w:after'),'0'); e.set(qn('w:line'),'312' if natural else '480'); e.set(qn('w:lineRule'),'auto' if natural else 'exact'); pr.append(e)
    e=OxmlElement('w:ind'); e.set(qn('w:firstLineChars'),'0' if role in ('image','caption','title','table_caption') else '200'); pr.append(e)
    e=OxmlElement('w:jc'); e.set(qn('w:val'),'center' if role in ('image','caption','title','table_caption') else 'both'); pr.append(e)
    if role in ('image','heading','title','table_caption','results_analysis_heading','benefit_intro','auxiliary_label'):
        e=OxmlElement('w:keepNext'); pr.append(e)
    for r in p.xpath('.//w:r'):
        rp=r.get_or_add_rPr()
        for tag in ['rFonts','sz','szCs','color','highlight','b','bCs']:
            e=rp.find(qn('w:'+tag))
            if e is not None: rp.remove(e)
        e=OxmlElement('w:rFonts'); e.set(qn('w:eastAsia'),'宋体'); e.set(qn('w:ascii'),'Times New Roman'); e.set(qn('w:hAnsi'),'Times New Roman'); rp.append(e)
        for tag in ['sz','szCs']:
            e=OxmlElement('w:'+tag); e.set(qn('w:val'),'24'); rp.append(e)
        e=OxmlElement('w:color'); e.set(qn('w:val'),'000000'); rp.append(e)
        for tag in ('b','bCs'):
            flag=OxmlElement('w:'+tag);flag.set(qn('w:val'),'1' if role in ('heading','title') else '0');rp.append(flag)

def replace_text(node,old,new,required=False):
    nodes=node.xpath('.//w:t'); full=''.join(n.text or '' for n in nodes)
    if required and old not in full: raise ValueError('Replacement target absent: '+old)
    # Rebuild only affected text spans; untouched run formatting and superscripts remain.
    while old in full:
        start=full.index(old); end=start+len(old); pos=0; inserted=False
        for n in nodes:
            s=n.text or ''; a=pos; b=pos+len(s); pos=b
            if b<=start or a>=end: continue
            left=s[:max(0,start-a)]; right=s[max(0,end-a):] if end<b else ''
            n.text=left+(new if not inserted else '')+right; inserted=True
        full=''.join(n.text or '' for n in nodes)
        if old in new: break

def slice_p(node,start,end):
    n=deepcopy(node); full=text(n)
    if not 0<=start<=end<=len(full): raise ValueError('Invalid source text slice')
    if (start or end!=len(full)) and n.xpath('.//m:oMath | .//w:drawing | .//w:pict | .//w:object | .//w:sym | .//w:fldChar | .//w:fldSimple | .//w:instrText | .//w:footnoteReference | .//w:endnoteReference | .//w:tab | .//w:br | .//w:cr'):
        raise ValueError('Cannot slice a paragraph containing math/image/OLE/field/tab/break; import whole rich object or author explicit runs')
    pos=0
    for t in n.xpath('.//w:t'):
        s=t.text or ''; a=pos; b=a+len(s); pos=b
        t.text=s[max(0,start-a):max(0,min(len(s),end-a))] if b>start and a<end else ''
    return n

def claim_layout(nodes):
    """Normalize real w:p boundaries, keeping source runs and native scripts."""
    from bio_claims import FEATURE, paragraph_spans
    groups=[];changes=[];result=[]
    for node in nodes:
        if ET.QName(node).localname!='p':raise ValueError('Claims must use text paragraphs, not tables')
        if node.xpath('.//w:br | .//w:cr | .//w:tab'):
            raise ValueError('Claims require natural paragraphs; explicit breaks/tabs need native editing')
        value=text(node)
        if not value.strip() and node.xpath('.//m:oMath | .//w:drawing | .//w:pict | .//w:object | .//w:sym | .//w:fldChar | .//w:fldSimple | .//w:instrText | .//w:footnoteReference | .//w:endnoteReference'):
            raise ValueError('Textless rich claim object needs source-bound native placement; cannot rebuild by text offsets')
        if re.match(r'^\s*\d+、',value):groups.append([])
        if not groups:raise ValueError('Claim text before the first numbered item')
        if value.strip() or node.xpath('.//m:oMath | .//w:drawing | .//w:object'):groups[-1].append(node)
    for group in groups:
        values=[text(n) for n in group];flat=''.join(values)
        head=flat.split(FEATURE,1)[0]
        kind='dependent' if re.match(r'^\s*\d+、根据权利要求',head) else 'independent'
        spans=paragraph_spans(flat,kind);expected=[flat[a:b] for a,b in spans]
        if values==expected:normalized=group
        else:
            normalized=[]
            for start,end in spans:
                paragraph=OxmlElement('w:p');pos=0
                for node,value in zip(group,values):
                    stop=pos+len(value)
                    if stop>start and pos<end:
                        part=slice_p(node,max(0,start-pos),min(len(value),end-pos))
                        for child in list(part):
                            if ET.QName(child).localname!='pPr':paragraph.append(child)
                    pos=stop
                normalized.append(paragraph)
            assert ''.join(text(n) for n in normalized)==flat,'Claim layout must not change text'
            changes.append({'number':int(re.match(r'^\s*(\d+)、',flat)[1]),
                            'before_paragraphs':len(group),'after_paragraphs':len(normalized),
                            'text_changed':False})
        for node in normalized:format_p(node,'body','claims')
        result.extend(normalized)
    return result,changes

def build(evidence_path,plan_path,output):
    evidence_path=Path(evidence_path); plan_path=Path(plan_path); output=Path(output)
    if output.exists(): raise ValueError('Refusing to overwrite output')
    e=read(evidence_path); plan=read(plan_path)
    if plan.get('source_sha256')!=e['source']['sha256'] or sha(e['source']['path'])!=e['source']['sha256']: raise ValueError('Source identity mismatch')
    srcpath=Path(e['working_source']['path'])
    if sha(srcpath)!=e['working_source']['sha256']: raise ValueError('Working DOCX changed')
    plan,plan_changes=normalize_plan(plan,e)
    if not plan.get('title') or not plan.get('claim_map'): raise ValueError('Title and semantic claim_map required')
    claims=plan['claim_map']
    from bio_claims import collect, validate_plan
    claim_paragraphs=[];claim_layout_changes=[]
    src=Document(srcpath); dest=Document(SKILL/'assets/patent-template.docx'); imported={}; trace=[]
    for name in ('BioFigureAnalysis','BioResultsHeading','BioTechnicalSummary'):
        if name not in dest.styles:
            style=dest.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.base_style=dest.styles['Normal']
    def import_part(part):
        key=str(part.partname)
        if key in imported: return imported[key]
        folder='media' if part.content_type.startswith('image/') else ('embeddings' if '/embeddings/' in key else 'imported')
        name='/word/'+folder+'/'+uuid.uuid4().hex+'_'+PurePosixPath(key).name
        obj=Part(PackURI(name),part.content_type,part.blob,dest.part.package); imported[key]=obj
        replacements={}
        for rid,rel in part.rels.items():
            newrid=obj.relate_to(rel.target_ref if rel.is_external else import_part(rel.target_part),rel.reltype,rel.is_external); replacements[rid]=newrid
        if 'xml' in part.content_type:
            try:
                root=ET.fromstring(part.blob)
                for node in root.iter():
                    for attr,val in list(node.attrib.items()):
                        if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and val in replacements: node.set(attr,replacements[val])
                obj._blob=ET.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
            except ET.XMLSyntaxError: pass
        return obj
    def remap(n):
        for node in n.iter():
            for attr,rid in list(node.attrib.items()):
                if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    rel=src.part.rels[rid]; target=rel.target_ref if rel.is_external else import_part(rel.target_part)
                    node.set(attr,dest.part.relate_to(target,rel.reltype,rel.is_external))
        for sect in n.xpath('.//w:sectPr'): sect.getparent().remove(sect)
        for bad in n.xpath('.//w:bookmarkStart | .//w:bookmarkEnd | .//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference'):
            bad.getparent().remove(bad)
        # Avoid duplicate drawing identities when the main image is reused.
        for dp in n.xpath('.//wp:docPr'): dp.set('id',str(len(trace)*100+len(imported)+1))
        return n
    def make(block,slot):
        role=block.get('role','body')
        if 'source_body_index' in block:
            idx=block['source_body_index']; original=src.element.body[idx]
            if ET.QName(original).localname not in ('p','tbl'): raise ValueError('Only paragraph/table sources may be imported')
            n=slice_p(original,*block['slice']) if 'slice' in block else deepcopy(original)
            n=remap(n)
            if block.get('prefix'):
                r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=block['prefix']; r.append(t)
                n.insert(1 if n.pPr is not None else 0,r)
            for old,new in block.get('replace',[]): replace_text(n,old,new,required=True)
            trace.append({'slot':slot,'body_index':idx,'slice':block.get('slice'),'replacements':block.get('replace',[]),'prefix':block.get('prefix','')})
        else:
            n=OxmlElement('w:p')
            for item in block.get('runs',[{'text':block.get('text','')}]):
                r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=item['text']; r.append(t)
                if item.get('script'):
                    if item['script'] not in ('subscript','superscript'): raise ValueError('Invalid script')
                    rp=OxmlElement('w:rPr'); va=OxmlElement('w:vertAlign'); va.set(qn('w:val'),item['script']); rp.append(va); r.insert(0,rp)
                n.append(r)
            trace.append({'slot':slot,'authored_text':text(n),'source_body_indices':block.get('source_body_indices',[])})
            role=paragraph_role(role,text(n),slot,plan['title'])
            if text(n).strip() and role not in ('heading','title','caption') and not block.get('source_body_indices'): raise ValueError('Authored technical prose needs source locators')
        role=paragraph_role(role,text(n),slot,plan['title'])
        for old,new in REPLACE: replace_text(n,old,new)
        if slot!='claims': replace_text(n,'，其特征在于，','，'); replace_text(n,'其特征在于，','')
        if ET.QName(n).localname=='p':
            if slot=='description' and (role=='technical_summary' or SUMMARY.match(text(n))):
                trace[-1]['summary_normalization']=summary_format(n,block,replace_text)
            format_p(n,role,slot)
            if role=='image':
                # Source aspect ratio is preserved; never upscale or crop.
                for ext in n.xpath('.//wp:extent'):
                    cx=int(ext.get('cx')); cy=int(ext.get('cy')); scale=min(1,5800000/max(cx,1),7200000/max(cy,1)); nx=int(cx*scale); ny=int(cy*scale); ext.set('cx',str(nx)); ext.set('cy',str(ny))
                    for x in n.xpath('.//a:xfrm/a:ext'): x.set('cx',str(nx)); x.set('cy',str(ny))
        else:
            for p in n.xpath('.//w:p'): format_p(p,'table')
            trace[-1]['table_layout']=table_format(n,block.get('table_header_rows'))
        return n
    for slot in SLOTS:
        marker=next(p for p in dest.paragraphs if p.text=='BIO_SLOT_'+slot.upper())
        blocks=plan.get('sections',{}).get(slot)
        if blocks is None: raise ValueError('Missing section: '+slot)
        if not blocks and slot not in ('abstract_figure','figures'): raise ValueError('Required section empty')
        boundary=marker._p.pPr.find(qn('w:sectPr')) if marker._p.pPr is not None else None
        last=None
        nodes=[make(block,slot) for block in blocks]
        if slot=='claims':nodes,claim_layout_changes=claim_layout(nodes)
        for current in nodes:
            if slot=='claims':claim_paragraphs.append({'text':text(current),'p':len(claim_paragraphs)})
            if last is not None and ET.QName(last).localname=='tbl' and ET.QName(current).localname=='tbl':
                # Word/WPS can merge adjacent w:tbl into one table and reuse the
                # preceding table's repeated header despite different schemas.
                spacer=OxmlElement('w:p');pr=spacer.get_or_add_pPr()
                spacing=OxmlElement('w:spacing');spacing.set(qn('w:line'),'20');spacing.set(qn('w:lineRule'),'exact');spacing.set(qn('w:before'),'0');spacing.set(qn('w:after'),'0');pr.append(spacing)
                snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');pr.append(snap)
                marker._p.addprevious(spacer)
                trace.append({'slot':slot,'layout_only':'separator_between_distinct_source_tables','technical_content_added':False})
            last=current; marker._p.addprevious(last)
        if boundary is not None:
            if last is None or ET.QName(last).localname!='p':
                last=OxmlElement('w:p'); marker._p.addprevious(last)
            last.get_or_add_pPr().append(deepcopy(boundary))
        marker._p.getparent().remove(marker._p)
    claim_source={'source':e['source'],'body':[{'body_index':i,'text':text(n)} for i,n in enumerate(src.element.body) if ET.QName(n).localname!='sectPr']}
    claim_validation=validate_plan(collect(claim_paragraphs),plan,claim_source)
    from bio_spacing import forbidden
    from bio_presentation import normalize_figure_description_endings
    figure_endings=normalize_figure_description_endings(dest.element)
    if any(forbidden(c) for c in plan['title']):raise ValueError('Formal title needs source-supported no-space wording before freezing the plan')
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='bio-no-spaces-',dir=output.parent) as tempdir:
        pending=Path(tempdir)/'pending.docx';dest.save(pending)
        from bio_package import normalize
        formatted=normalize(pending)
        scripts=formatted['native_scripts'];background=formatted['background'];spacing=formatted['spacing']
        with output.open('xb') as stream:stream.write(pending.read_bytes())
    write(output.with_suffix('.trace.json'),{'schema_version':1,'artifact_sha256':sha(output),'source_sha256':e['source']['sha256'],'plan_sha256':sha(plan_path),'plan_normalizations':plan_changes,'claim_layout_normalizations':claim_layout_changes,'title':plan['title'],'claim_map':claims,'claim_validation':claim_validation,'blocks':trace,'spacing':spacing,'background':background,'native_scripts':scripts,'figure_description_endings':figure_endings,'status':'draft_not_accepted'})
    return output
def check_generated(output,evidence_path,plan_path):
    """Run the existing checks once and publish the writer's reusable cache."""
    from bio_check import inspect
    import inspect_docx
    output=Path(output); preflight=output.with_suffix('.preflight.json')
    report=inspect(output,read(plan_path),read(evidence_path));write(preflight,report)
    if report['errors']:
        raise ValueError('Generated draft failed mechanical preflight; see '+str(preflight))
    inspection=inspect_docx.run_inspection(output,output.parent/'inspection-cache','draft',True,True)
    if sha(output)!=report['artifact_sha256']:
        raise ValueError('Generated draft changed between mechanical checks')
    return {'status':'draft_not_accepted','artifact':str(output.resolve()),'preflight':str(preflight.resolve()),'inspection':inspection}

def main():
    p=argparse.ArgumentParser(); p.add_argument('--evidence',required=True); p.add_argument('--plan',required=True); p.add_argument('--out',required=True)
    p.add_argument('--check',action='store_true',help='Run mechanical preflight and create the normal inspection-cache; independent review is still required')
    a=p.parse_args();output=build(a.evidence,a.plan,a.out)
    print(__import__('json').dumps(check_generated(output,a.evidence,a.plan),ensure_ascii=False) if a.check else output)
if __name__=='__main__':main()
