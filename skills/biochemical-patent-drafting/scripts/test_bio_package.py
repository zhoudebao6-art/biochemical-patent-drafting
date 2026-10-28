"""Transactional formatting and unchanged validation boundaries; no model calls."""
from pathlib import Path
import hashlib
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
import bio_package
import bio_background
import bio_spacing
import bio_scripts
import json
import stage_bundle
import inspect_docx
from collections import Counter


def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()


class PackageTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name)/'sample.docx'

    def fixture(self, ambiguous=False):
        doc = Document(); p = doc.add_paragraph('Na₂CO₃含量为0.5 g；10⁻⁵。')
        props = p.runs[0]._r.get_or_add_rPr()
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'FFFFFF'); props.append(shade)
        background = OxmlElement('w:background'); background.set(qn('w:color'), 'CCFFCC'); doc.element.insert(0, background)
        doc.sections[0].header.paragraphs[0].text = '说 明 书'
        if ambiguous: doc.add_paragraph('retain separate words')
        doc.save(self.path)
        # Non-XML data must be byte-identical after normalization.
        with zipfile.ZipFile(self.path, 'a') as package: package.writestr('word/media/unchanged.bin', b'original-scientific-object')

    def inspection_fixture(self):
        """Synthetic display objects only; no user case or experimental data."""
        self.fixture()
        doc=Document(self.path)
        doc.sections[0].footer.paragraphs[0].text='页脚示例'
        doc.add_section()
        # Repeat an explicit reference to the same header across two sections.
        ref=doc.sections[0]._sectPr.find(qn('w:headerReference'))
        from copy import deepcopy
        doc.sections[-1]._sectPr.insert(0,deepcopy(ref))
        p=doc.add_paragraph('检查颜色继承'); p.runs[0].font.color.rgb=__import__('docx.shared',fromlist=['RGBColor']).RGBColor(255,0,0)
        doc.save(self.path)
        with zipfile.ZipFile(self.path,'a') as z:
            w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
            for name,tag in [('footnotes','footnote'),('comments','comment')]:
                z.writestr('word/'+name+'.xml',f'<w:{name} xmlns:w="{w}"><w:{tag} w:id="1"><w:p><w:r><w:t>{tag}测试</w:t></w:r></w:p></w:{tag}></w:{name}>')
            z.writestr('word/charts/chart1.xml','<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:p><a:r><a:t>图表标签 10⁻⁵</a:t></a:r></a:p></c:chartSpace>')
            z.writestr('word/drawings/test.vml','<xml xmlns:v="urn:schemas-microsoft-com:vml"><v:shape><v:textpath string="测试 标签"/></v:shape></xml>')

    def test_standard_inspector_parses_each_xml_member_once(self):
        self.inspection_fixture();original=sha(self.path)
        with zipfile.ZipFile(self.path) as z:
            xml={hashlib.sha256(z.read(n)).hexdigest():n for n in z.namelist() if n.endswith(('.xml','.rels','.vml'))}
        counts=Counter();parse=inspect_docx.ET.fromstring
        def counted(payload,*args,**kwargs):
            data=payload.encode() if isinstance(payload,str) else payload
            name=xml.get(hashlib.sha256(data).hexdigest())
            if name: counts[name]+=1
            return parse(payload,*args,**kwargs)
        with patch.object(inspect_docx.ET,'fromstring',side_effect=counted):
            report=inspect_docx.inspect(self.path)
        self.assertGreaterEqual(len(counts),10)
        self.assertTrue(all(count==1 for count in counts.values()),dict(counts))
        self.assertEqual(sha(self.path),original)
        scans=bio_package.scan_docx(self.path)
        for mode in ('background','native_scripts','spacing'):
            self.assertEqual(report['biochemical_'+mode+'_evidence'],scans[mode])
        self.assertTrue(report['text_color_evidence']['review_runs'])

    def test_xml_cache_does_not_cross_inspections(self):
        self.inspection_fixture();first=inspect_docx.inspect(self.path)
        with zipfile.ZipFile(self.path) as z: entries=[(i,z.read(i.filename)) for i in z.infolist()]
        with zipfile.ZipFile(self.path,'w') as z:
            for i,data in entries:
                z.writestr(i,data.replace('页脚示例'.encode(), '已修改页脚'.encode()) if i.filename.startswith('word/footer') else data)
        second=inspect_docx.inspect(self.path)
        self.assertNotEqual(first['sha256'],second['sha256'])
        self.assertTrue(any('已修改页脚' in p['text'] for part in second['parts'] for p in part['paragraphs']))

    def test_xml_reader_is_local_to_one_inspection(self):
        self.inspection_fixture()
        with zipfile.ZipFile(self.path) as z:
            reader=inspect_docx.xml_reader(z)
            with patch.object(z,'read',wraps=z.read) as read:
                body=reader('word/document.xml')
                self.assertIs(reader('word/document.xml'),body)
                self.assertEqual(read.call_count,1)
            self.assertIsNot(inspect_docx.xml_reader(z)('word/document.xml'),body)

    def test_single_commit_and_final_hash_preserve_objects_and_native_scripts(self):
        self.fixture()
        replace = bio_package.os.replace
        with patch.object(bio_package.os, 'replace', wraps=replace) as commit:
            result = bio_package.normalize(self.path)
        self.assertEqual(commit.call_count, 1)
        doc = Document(self.path); self.assertEqual(doc.paragraphs[0].text, 'Na2CO3含量为0.5g；10-5。')
        self.assertEqual([r.text for r in doc.paragraphs[0].runs if r.font.subscript], ['2', '3'])
        self.assertEqual([r.text for r in doc.paragraphs[0].runs if r.font.superscript], ['-5'])
        for mode in ('background', 'native_scripts'):
            self.assertEqual(result[mode]['evidence']['artifact_sha256'], sha(self.path))
            self.assertEqual(result[mode]['evidence']['issues'], [])
        with zipfile.ZipFile(self.path) as package: self.assertEqual(package.read('word/media/unchanged.bin'), b'original-scientific-object')
        stable = sha(self.path); bio_package.normalize(self.path); self.assertEqual(sha(self.path), stable)

    def test_late_ambiguity_leaves_entire_file_untouched(self):
        self.fixture(ambiguous=True); original = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, 'source-bound rewording'):
            bio_package.normalize(self.path)
        self.assertEqual(self.path.read_bytes(), original)
        self.assertTrue(bio_background.scan_docx(self.path)['issues'])
        self.assertTrue(bio_scripts.scan_docx(self.path)['issues'])

    def test_shared_scan_parses_each_xml_once_and_keeps_all_hits(self):
        self.fixture(); parse = ET.fromstring
        with zipfile.ZipFile(self.path) as package:
            expected = len(bio_spacing.selected_parts(package.namelist()))
            with patch.object(bio_package.ET, 'fromstring', wraps=parse) as reader:
                scans = bio_package.scan_package(package, sha(self.path))
            self.assertEqual(reader.call_count, expected)
        for mode, module, key in [('spacing', bio_spacing, 'hits'), ('background', bio_background, 'issues'), ('native_scripts', bio_scripts, 'issues')]:
            self.assertEqual(scans[mode][key], module.scan_docx(self.path)[key])

    def test_individual_domain_entrypoints_preserve_other_domains(self):
        self.fixture(); bio_background.normalize(self.path)
        self.assertFalse(bio_background.scan_docx(self.path)['issues'])
        self.assertTrue(bio_scripts.scan_docx(self.path)['issues'])
        self.assertTrue(bio_spacing.scan_docx(self.path)['hits'])
        bio_scripts.normalize(self.path); self.assertFalse(bio_scripts.scan_docx(self.path)['issues'])
        self.assertTrue(bio_spacing.scan_docx(self.path)['hits'])
        bio_spacing.normalize(self.path); self.assertFalse(bio_spacing.scan_docx(self.path)['hits'])

    def test_compact_packet_binds_exact_exceptions_without_dropping_rules(self):
        skill = Path(__file__).resolve().parents[1]
        rules = json.loads((skill/'references/rules.json').read_text(encoding='utf-8'))
        for stage in stage_bundle.STAGES:
            packet, coverage = stage_bundle.build_packet(skill, rules, stage)
            shared = {rule_id: row for row in coverage['shared_exceptions'] for rule_id in row['rule_ids']}
            for rule in rules['rules']:
                if rule['status'] != 'active' or rule['id'] not in coverage['included_rule_ids']: continue
                row = shared.get(rule['id'])
                if row:
                    self.assertEqual(row['exceptions'], rule['exceptions'])
                    self.assertIn(stage_bundle.rule_block(rule, row['id']), packet)
                    self.assertEqual(packet.count(stage_bundle.canonical(row['exceptions'])), 1)
                else: self.assertIn(stage_bundle.rule_block(rule), packet)


if __name__ == '__main__': unittest.main()
