"""Structural fixtures only; these are not technical or experimental examples."""
import copy
import unittest
from bio_claims import candidates,collect,analyze,source_check,validate_plan,dependency_check


def fixture(groups=3):
    source=['工作原理','原理说明',
            '为了解决上述技术问题，本发明提供了一种材料A，其特征在于，包括组分甲。',
            '进一步地，在上述技术方案中，所述组分甲为颗粒。']
    texts=['1、一种材料A，其特征在于，包括组分甲。',
           '2、根据权利要求1所述的材料A，其特征在于，所述组分甲为颗粒。']
    mapping=[{'number':1,'parent':None,'object':'一种材料A','source_body_indices':[2]},
             {'number':2,'parent':1,'object':'材料A','source_body_indices':[3],'dependency_reason':'继承材料A组分限定'}]
    if groups>=2:
        source += ['本方案还提供一种材料A的制备方法，其特征在于，包括以下步骤：',
                   'S1、混合组分甲；','S2、分离并获得所述材料A。',
                   '具体实施中，解释不生成新权项。',
                   '进一步地，在上述技术方案中，步骤S1使用搅拌器。']
        texts += ['3、一种权利要求1至2任一项所述的材料A的制备方法，其特征在于，包括以下步骤：',
                  'S1、混合组分甲；','S2、分离并获得所述材料A。',
                  '4、根据权利要求3所述的材料A的制备方法，其特征在于，步骤S1使用搅拌器。']
        mapping += [{'number':3,'kind':'independent','parent':None,'group':3,'object':'一种材料A的制备方法',
                     'head':'一种权利要求1至2任一项所述的材料A的制备方法','reference_claims':[1,2],
                     'source_body_indices':[4,5,6]},
                    {'number':4,'parent':3,'group':3,'object':'材料A的制备方法','source_body_indices':[8],
                     'dependency_reason':'继承方法组S1混合操作'}]
    if groups>=3:
        source += ['本发明还提供了材料A的应用，其特征在于，用于示例检测。',
                   '进一步地，在上述技术方案中，所述示例检测使用标签。']
        texts += ['5、一种权利要求1至2任一项所述的材料A的应用，其特征在于，用于示例检测。',
                  '6、根据权利要求5所述的材料A的应用，其特征在于，所述示例检测使用标签。']
        mapping += [{'number':5,'kind':'independent','parent':None,'object':'一种材料A的应用',
                     'head':'一种权利要求1至2任一项所述的材料A的应用','reference_claims':[1,2],'source_body_indices':[9]},
                    {'number':6,'parent':5,'object':'材料A的应用','source_body_indices':[10],
                     'dependency_reason':'继承应用组示例检测对象'}]
    needs={2:(1,2,'组分甲','所述组分甲为颗粒'),4:(3,5,'混合组分甲','步骤S1使用搅拌器'),6:(5,9,'用于示例检测','所述示例检测使用标签')}
    for m in mapping:
        if m['number'] in needs:
            parent,idx,quote,current=needs[m['number']]
            m['dependency_basis']=[{'claim':parent,'body_index':idx,'quote':quote,'current_quote':current}]
    claims=collect([{'p':i,'text':t} for i,t in enumerate(texts)])
    # These mutable string fixtures do not represent a frozen Word inventory.
    for c in claims:c.pop('paragraph_texts')
    return claims,{'claim_map':mapping}, {
        'source':{'sha256':'fixture_only'},'body':[{'body_index':i,'text':t} for i,t in enumerate(source)]}


class ClaimTests(unittest.TestCase):
    def test_internal_where_clause_does_not_split_a_numbered_step(self):
        from bio_claims import paragraph_spans
        value='1、一种方法，其特征在于，包括以下步骤：S1、混合；其中，混合过程中保持密闭；S2、分离。'
        got=[value[a:b] for a,b in paragraph_spans(value,'independent')]
        self.assertEqual(len(got),3)
        self.assertEqual(got[1],'S1、混合；其中，混合过程中保持密闭；')

    def test_final_where_clause_does_not_invent_a_numbered_step_boundary(self):
        from bio_claims import paragraph_spans
        value='1、一种方法，其特征在于，包括以下步骤：S1、混合；S2、分离；其中，步骤S2的分离过程中保持密闭。'
        got=[value[a:b] for a,b in paragraph_spans(value,'independent')]
        self.assertEqual(len(got),3)
        self.assertEqual(got[2],'S2、分离；其中，步骤S2的分离过程中保持密闭。')

    def test_xml_break_metadata_is_not_ignored(self):
        claims=collect([{'text':'1、一种方法，其特征在于，包括甲。','p':0},
                        {'text':'2、根据权利要求1所述的方法，其特征在于，混合；分离。','p':1,'manual_breaks':1}])
        self.assertTrue(any('paragraph layout' in s for s in analyze(claims)['issues']))
        from docx import Document
        from bio_author import claim_layout
        d=Document();p=d.add_paragraph(claims[0]['text']);p.add_run().add_break()
        with self.assertRaisesRegex(ValueError,'explicit breaks'):claim_layout([p._p])

    def test_textless_rich_object_is_not_silently_dropped(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from bio_author import claim_layout
        for kind in ('w:drawing','w:pict','w:object','m:oMath','w:sym'):
            d=Document();p=d.add_paragraph('1、一种方法，其特征在于，包括以下步骤：S1、混合；')
            image=d.add_paragraph();image._p.append(OxmlElement(kind));q=d.add_paragraph('S2、分离。')
            with self.subTest(kind=kind),self.assertRaisesRegex(ValueError,'Textless rich'):claim_layout([p._p,image._p,q._p])

    def test_rich_object_in_text_is_not_duplicated_when_splitting(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from bio_author import claim_layout
        for kind in ('w:sym','w:pict','w:fldSimple','w:instrText','w:footnoteReference','w:endnoteReference'):
            d=Document();p=d.add_paragraph('1、一种方法，其特征在于，包括以下步骤：S1、混合；S2、分离。')
            p._p.append(OxmlElement(kind));original=p._p.xml
            with self.subTest(kind=kind),self.assertRaisesRegex(ValueError,'Cannot slice'):claim_layout([p._p])
            self.assertEqual(p._p.xml,original)

    def test_matching_plan_and_text_cannot_replace_source_object(self):
        claims,plan,e=fixture(1)
        for c in claims:c['text']=c['text'].replace('材料A','另一种无来源材料')
        for m in plan['claim_map']:m['object']=m['object'].replace('材料A','另一种无来源材料')
        with self.assertRaisesRegex(ValueError,'object differs from source'):validate_plan(claims,plan,e)

    def test_candidate_order_is_not_only_a_set(self):
        claims,plan,e=self.branch_fixture()
        # Independent sibling refinements can each pass dependency checking,
        # but exchanging their source order is still not the supplied scheme.
        claims[1],claims[3]=claims[3],claims[1]
        plan['claim_map'][1],plan['claim_map'][3]=plan['claim_map'][3],plan['claim_map'][1]
        for n in (2,4):
            claims[n-1]['number']=n;claims[n-1]['text']=str(n)+claims[n-1]['text'][1:]
            plan['claim_map'][n-1]['number']=n
        self.assertTrue(any('order changed' in s for s in source_check(claims,plan['claim_map'],e)['issues']))

    def test_independent_method_cannot_remain_one_paragraph(self):
        claims,plan,e=fixture(2);claims[2]['text']=claims[2]['text'].replace('\n','')
        with self.assertRaisesRegex(ValueError,'paragraph layout'):validate_plan(claims,plan,e)

    def test_dependent_procedure_stays_one_paragraph(self):
        text='2、根据权利要求1所述的方法，其特征在于，前体按以下步骤获得：混合；分离。'
        claims=[{'number':1,'text':'1、一种方法，其特征在于，包括前体。'}, {'number':2,'text':text}]
        self.assertFalse(analyze(claims)['issues'])
        claims[1]['text']=text.replace('：','：\n').replace('；','；\n')
        self.assertTrue(any('paragraph layout' in s for s in analyze(claims)['issues']))

    def test_physical_paragraphs_not_soft_line_breaks(self):
        claims,_,_=fixture(2)
        c=collect([{'p':0,'text':claims[2]['text']}])[0]
        from bio_claims import paragraph_check
        self.assertFalse(paragraph_check(c,'independent')['valid'])
        self.assertTrue(paragraph_check(claims[2],'independent')['valid'])

    def test_numbered_method_keeps_internal_semicolons(self):
        from bio_claims import paragraph_spans
        text='1、一种方法，其特征在于，包括如下步骤：S1、混合；S2、载钙；制酸液；S3、接触；过滤。'
        actual=[text[a:b] for a,b in paragraph_spans(text,'independent')]
        self.assertEqual(len(actual),4)
        self.assertEqual(actual[2],'S2、载钙；制酸液；')
        self.assertEqual(actual[3],'S3、接触；过滤。')

    def test_unnumbered_method_preserves_tail_and_nested_punctuation(self):
        from bio_claims import paragraph_spans
        text='1、一种方法，其特征在于，包括以下步骤：混合（先甲；后乙）；析晶；其中，采用无水介质。'
        actual=[text[a:b] for a,b in paragraph_spans(text,'independent')]
        self.assertEqual(actual[1: ],['混合（先甲；后乙）；','析晶；','其中，采用无水介质。'])
        self.assertNotIn('S1',''.join(actual))

    def test_unrecorded_transition_not_allowed_inside_dependent(self):
        claims,plan,e=fixture(1)
        e['body'][3]['text']+='再分散；混合。'
        claims[1]['text']+='再分散；其中，混合。'
        with self.assertRaisesRegex(ValueError,'limitations changed'):validate_plan(claims,plan,e)

    def test_explicit_edit_cannot_change_units_or_chemical_tokens(self):
        for old,new in [('100mg/kg','100g/kg'),('CaCl2','NaCl2'),('0.5mol/L','0.5mmol/L')]:
            claims,plan,e=fixture(1)
            e['body'][3]['text']=e['body'][3]['text'].replace('颗粒',old+'颗粒')
            claims[1]['text']=claims[1]['text'].replace('颗粒',new+'颗粒')
            plan['claim_map'][1]['source_text_edits']=[{'old':old,'new':new,'reason':'不得放行的测试编辑'}]
            with self.subTest(old=old),self.assertRaisesRegex(ValueError,'unit/chemical'):validate_plan(claims,plan,e)

    def test_layout_writer_keeps_native_script_runs(self):
        from docx import Document
        from bio_author import claim_layout
        from bio_prepare import text
        d=Document();p=d.add_paragraph('1、一种方法，其特征在于，包括以下步骤：S1、加入Na')
        p.add_run('2').font.subscript=True;p.add_run('CO');p.add_run('3').font.subscript=True
        p.add_run('；S2、处理；分离。')
        original=p.text;nodes,changes=claim_layout([p._p])
        self.assertEqual(len(nodes),3);self.assertEqual(''.join(text(n) for n in nodes),original)
        self.assertEqual(len(nodes[1].xpath('.//w:vertAlign')),2)
        self.assertEqual(changes[0]['after_paragraphs'],3)
        # A second pass is stable, and does not accumulate empty split runs.
        second,changes=claim_layout(nodes);self.assertFalse(changes)
        self.assertEqual([n.xml for n in second],[n.xml for n in nodes])

    def test_layout_writer_merges_dependent_fragments(self):
        from docx import Document
        from bio_author import claim_layout
        from bio_prepare import text
        d=Document();parts=['1、一种方法，其特征在于，包括甲。',
            '2、根据权利要求1所述的方法，其特征在于，按以下步骤获得：','混合；','分离。']
        nodes,changes=claim_layout([d.add_paragraph(t)._p for t in parts])
        self.assertEqual(len(nodes),2);self.assertEqual(text(nodes[1]),''.join(parts[1:]))
        self.assertEqual(changes[0]['before_paragraphs'],3)

    def test_layout_writer_refuses_to_flatten_complex_objects(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from bio_author import claim_layout
        d=Document();p=d.add_paragraph('1、一种方法，其特征在于，包括以下步骤：S1、混合；S2、分离。')
        p._p.append(OxmlElement('m:oMath'))
        with self.assertRaisesRegex(ValueError,'Cannot slice'):claim_layout([p._p])

    def branch_fixture(self):
        source=['工作原理',
            '为了解决上述技术问题，本发明提供了一种材料A，其特征在于，包括组分甲和组分乙。',
            '进一步地，在上述技术方案中，所述组分甲为粒径1-5mm的颗粒。',
            '进一步地，在上述技术方案中，所述颗粒的粒径为2-3mm。',
            '进一步地，在上述技术方案中，所述组分乙为粉末。']
        parents=[None,1,2,1];texts=[];mapping=[]
        for n,parent in enumerate(parents,1):
            body=source[n].split('其特征在于，',1)[-1] if n==1 else source[n].split('在上述技术方案中，')[1]
            head='一种材料A' if parent is None else f'根据权利要求{parent}所述的材料A'
            texts.append({'number':n,'text':f'{n}、{head}，其特征在于，{body}'})
            m={'number':n,'parent':parent,'object':'材料A' if parent else '一种材料A','source_body_indices':[n]}
            if parent:
                quote={2:'包括组分甲和组分乙',3:'所述组分甲为粒径1-5mm的颗粒',4:'组分乙'}[n]
                m.update(dependency_reason={2:'限定权1组分甲的形态及粒径',3:'继承权2的颗粒形态及1-5mm范围并进一步收窄',4:'仅限定权1组分乙，不承接甲的形态或粒径条件'}[n],
                         dependency_basis=[{'claim':parent,'body_index':parent,'quote':quote,'current_quote':body}])
            mapping.append(m)
        return texts,{'claim_map':mapping},{'source':{'sha256':'fixture'},'body':[{'body_index':i,'text':t} for i,t in enumerate(source)]}

    def test_refinement_branches_and_complete_inherited_limits(self):
        claims,plan,e=self.branch_fixture();r=validate_plan(claims,plan,e)
        self.assertEqual([c['inheritance_paths'] for c in r['structure']['claims']],[[[1]],[[2,1]],[[3,2,1]],[[4,1]]])
        limits=r['source']['dependencies']['records'][1]['branches'][0]['inherited_limitations']
        self.assertEqual([x['claim'] for x in limits],[2,1]);self.assertIn('1-5mm',limits[0]['text'])

    def test_matching_wrong_plan_and_text_cannot_drop_required_parent(self):
        claims,plan,e=self.branch_fixture()
        claims[2]['text']=claims[2]['text'].replace('权利要求2','权利要求1');plan['claim_map'][2]['parent']=1
        self.assertFalse(analyze(claims,plan['claim_map'])['issues'])
        with self.assertRaisesRegex(ValueError,'does not inherit'):validate_plan(claims,plan,e)

    def test_previous_number_is_not_a_reason_to_add_restrictions(self):
        claims,plan,e=self.branch_fixture()
        claims[3]['text']=claims[3]['text'].replace('权利要求1','权利要求3');plan['claim_map'][3]['parent']=3
        with self.assertRaisesRegex(ValueError,'unexplained extra inheritance'):validate_plan(claims,plan,e)

    def test_noun_in_sibling_or_global_source_does_not_supply_inheritance(self):
        claims,plan,e=self.branch_fixture()
        e['body'].append({'body_index':5,'text':'具体实施中，前文出现所述颗粒，但不是权1的技术限定。'})
        claims[2]['text']=claims[2]['text'].replace('权利要求2','权利要求1');plan['claim_map'][2]['parent']=1
        with self.assertRaisesRegex(ValueError,'does not inherit'):validate_plan(claims,plan,e)

    def test_missing_or_source_mismatched_basis_rejected(self):
        for change in ('missing','wrong_location','only_reason','wrong_quote','wrong_current'):
            claims,plan,e=self.branch_fixture();m=plan['claim_map'][2]
            if change=='missing':m.pop('dependency_basis')
            elif change=='only_reason':m['dependency_basis']=[]
            elif change=='wrong_location':m['dependency_basis'][0]['body_index']=1
            elif change=='wrong_quote':m['dependency_basis'][0]['quote']='并不存在的技术条件'
            else:m['dependency_basis'][0]['current_quote']='另一项才有的限定'
            with self.subTest(change=change),self.assertRaises(ValueError):validate_plan(claims,plan,e)

    def test_alternative_paths_are_not_unioned(self):
        claims,plan,e=self.branch_fixture()
        claims[2]['text']=claims[2]['text'].replace('权利要求2','权利要求1至2任一项')
        m=plan['claim_map'][2];m.pop('parent');m.update(parents=[1,2],head='根据权利要求1至2任一项所述的材料A')
        with self.assertRaisesRegex(ValueError,'does not inherit'):validate_plan(claims,plan,e)

    def test_valid_alternative_checks_each_path(self):
        claims,plan,e=self.branch_fixture()
        claims[3]['text']=claims[3]['text'].replace('权利要求1','权利要求1至2任一项')
        m=plan['claim_map'][3];m.pop('parent');m.update(parents=[1,2],head='根据权利要求1至2任一项所述的材料A')
        m['dependency_basis'].append({'claim':2,'body_index':2,'quote':'所述组分甲为粒径1-5mm的颗粒','current_quote':'所述组分乙为粉末','via':2})
        m['dependency_reason']='乙粉末直接细化权1；择一的权2分支明确组合甲颗粒形态与范围'
        r=validate_plan(claims,plan,e)
        self.assertEqual([b['path'] for b in r['source']['dependencies']['records'][-1]['branches']],[[4,1],[4,2,1]])

    def test_alternative_cannot_hide_unexplained_extra_conditions(self):
        claims,plan,e=self.branch_fixture()
        claims[3]['text']=claims[3]['text'].replace('权利要求1','权利要求2至3任一项')
        m=plan['claim_map'][3];m.pop('parent');m.update(parents=[2,3],head='根据权利要求2至3任一项所述的材料A')
        result=dependency_check(claims,plan['claim_map'],e)
        self.assertEqual(sum('unexplained extra inheritance' in x for x in result['issues']),2)

    def test_wrong_object_cannot_be_reclassified_as_independent(self):
        claims,_,_=fixture(1);claims[1]['text']=claims[1]['text'].replace('所述的材料A','所述的材料A的制备方法')
        result=analyze(claims)
        self.assertTrue(any('protection object differs' in s for s in result['issues']))
        self.assertEqual(result['independent_count'],1)

    def test_one_two_three_objects_with_cross_references(self):
        for count in (1,2,3):
            claims,plan,e=fixture(count);r=validate_plan(claims,plan,e)
            self.assertEqual(r['structure']['independent_count'],count)
            if count>1:
                self.assertEqual(r['structure']['claims'][2]['parents'],[])
                self.assertEqual(r['structure']['claims'][2]['reference_claims'],[1,2])
                self.assertEqual(r['structure']['claims'][3]['parents'],[3])

    def test_candidate_has_continuations_but_excludes_explanation(self):
        _,_,e=fixture();c=candidates(e['body'])[2]
        self.assertEqual(c['source_body_indices'],[4,5,6])
        self.assertEqual(c['explanation_body_indices'],[7])
        self.assertEqual(c['group_candidate'],3)

    def test_actual_text_must_match_plan_not_merely_claim_count(self):
        claims,plan,e=fixture();claims[3]['text']=claims[3]['text'].replace('权利要求3','权利要求1')
        with self.assertRaises(ValueError):validate_plan(claims,plan,e)

    def test_title_cannot_replace_real_object(self):
        claims,plan,e=fixture(1)
        plan['claim_map'][1]['object']='材料A及其制备方法'
        with self.assertRaises(ValueError):validate_plan(claims,plan,e)

    def test_continuation_omission_and_duplication_are_rejected(self):
        claims,plan,e=fixture();bad=copy.deepcopy(claims)
        bad[2]['text']=bad[2]['text'].replace('S1、混合组分甲；\n','')
        self.assertTrue(source_check(bad,plan['claim_map'],e)['issues'])
        bad=copy.deepcopy(claims);bad[2]['text']+='S2、分离并获得所述材料A。'
        self.assertTrue(source_check(bad,plan['claim_map'],e)['issues'])
        plan['claim_map'][2]['source_body_indices'].remove(5)
        self.assertTrue(source_check(claims,plan['claim_map'],e)['issues'])

    def test_source_parameters_cannot_silently_change(self):
        claims,plan,e=fixture(1);e['body'][3]['text']=e['body'][3]['text'].replace('颗粒','10mg颗粒')
        with self.assertRaisesRegex(ValueError,'limitations'):validate_plan(claims,plan,e)

    def test_explicit_wording_edit_cannot_change_numbers(self):
        claims,plan,e=fixture(1)
        e['body'][3]['text']=e['body'][3]['text'].replace('颗粒','10mg颗粒')
        claims[1]['text']=claims[1]['text'].replace('颗粒','20mg颗粒')
        plan['claim_map'][1]['source_text_edits']=[{'old':'10mg','new':'20mg','reason':'不应放行的测试编辑'}]
        with self.assertRaisesRegex(ValueError,'numeric'):validate_plan(claims,plan,e)

    def test_missing_and_duplicate_objects_rejected(self):
        claims,plan,e=fixture()
        self.assertTrue(source_check(claims[:-2],plan['claim_map'][:-2],e)['issues'])
        plan['claim_map'][2]['source_body_indices']=[2]
        self.assertTrue(source_check(claims,plan['claim_map'],e)['issues'])

    def test_deep_combination_chain_is_retained(self):
        claims=[];mapping=[]
        for n in range(1,6):
            head='一种方法' if n==1 else f'根据权利要求{n-1}所述的方法'
            claims.append({'number':n,'text':f'{n}、{head}，其特征在于，条件{n}。'})
            mapping.append({'number':n,'parent':None if n==1 else n-1,'object':'一种方法' if n==1 else '方法',
                            'source_body_indices':[n],'dependency_reason':'保留前项条件组合'})
        result=analyze(claims,mapping);self.assertFalse(result['issues'])
        self.assertEqual([c['parents'] for c in result['claims']],[[],[1],[2],[3],[4]])

    def test_alternative_dependency_range(self):
        claims,plan,_=fixture(1)
        claims.append({'number':3,'text':'3、根据权利要求1至2任一项所述的材料A，其特征在于，带有标签。'})
        plan['claim_map'].append({'number':3,'parents':[1,2],'object':'材料A','head':'根据权利要求1至2任一项所述的材料A',
                                 'source_body_indices':[5],'dependency_reason':'每支均有相同材料对象，限制两支'})
        self.assertFalse(analyze(claims,plan['claim_map'])['issues'])
        claims[-1]['text']=claims[-1]['text'].replace('任一项','')
        self.assertTrue(analyze(claims)['issues'])

    def test_unparsed_or_future_references_rejected(self):
        for head in ('根据权利要求9所述的材料A','根据权利要求一所述的材料A','根据权利要求3至1任一项所述的材料A'):
            claims,_,_=fixture(1);claims[1]['text']='2、'+head+'，其特征在于，带有标签。'
            self.assertTrue(analyze(claims)['issues'])

    def test_wording_edit_is_explicit_and_remains_reviewable(self):
        claims,plan,e=fixture(1);claims[0]['text']=claims[0]['text'].replace('包括组分甲','包含组分甲')
        with self.assertRaises(ValueError):validate_plan(claims,plan,e)
        plan['claim_map'][0]['source_text_edits']=[{'old':'包括组分甲','new':'包含组分甲','reason':'同义引导，不改组分'}]
        result=validate_plan(claims,plan,e)
        self.assertTrue(result['source']['records'][0]['semantic_review_required'])


if __name__=='__main__':unittest.main()
