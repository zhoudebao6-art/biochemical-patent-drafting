"""Observable no-fill regressions; preserve technical objects and source bytes."""
import copy,unittest,zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET
from bio_common import sha,read
from bio_prepare import prepare
from bio_background import normalize,scan_docx
import test_biochemical as fixtures
import inspect_docx,bio_contract,batch_tool

def shade(parent,tag='shd'):
    n=OxmlElement('w:'+tag)
    if tag=='shd':n.set(qn('w:val'),'clear');n.set(qn('w:fill'),'CCFFCC');n.set(qn('w:themeFill'),'accent1')
    elif tag=='background':n.set(qn('w:color'),'CCFFCC')
    else:n.set(qn('w:val'),'yellow')
    parent.append(n)

def content(path):
    with zipfile.ZipFile(path) as z:
        root=ET.fromstring(z.read('word/document.xml'))
        return {'text':root.xpath('//w:t/text()',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}),
                'math':[ET.tostring(n,method='c14n') for n in root.xpath('//m:oMath',namespaces={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'})],
                'media':{n:z.read(n) for n in z.namelist() if n.startswith('word/media/')},
                'scripts':root.xpath('//w:vertAlign/@w:val',namespaces={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})}

class BackgroundTests(unittest.TestCase):
    def fixture(self):
        f=fixtures.BiochemicalBehaviorTests();f.setUp();self.addCleanup(f.doCleanups);return f
    def colored(self,f):
        out=f.build();d=Document(out)
        shade(d.element,'background')
        shade(d.paragraphs[0]._p.get_or_add_pPr());shade(d.paragraphs[0].runs[0]._r.get_or_add_rPr(),'highlight')
        shade(d.tables[0].cell(0,0)._tc.get_or_add_tcPr())
        shade(d.styles['Normal'].element.get_or_add_rPr())
        shade(d.sections[0].header.paragraphs[0]._p.get_or_add_pPr())
        styles=d.styles.element;table_style=OxmlElement('w:style');table_style.set(qn('w:type'),'table');table_style.set(qn('w:styleId'),'ColoredTable')
        conditional=OxmlElement('w:tblStylePr');conditional.set(qn('w:type'),'firstRow');pr=OxmlElement('w:tcPr');shade(pr);conditional.append(pr);table_style.append(conditional);styles.append(table_style)
        num=d.part.numbering_part.element;abstract=num.find(qn('w:abstractNum'));lvl=abstract.find(qn('w:lvl'));rp=OxmlElement('w:rPr');shade(rp,'highlight');lvl.append(rp)
        path=f.root/'colored.docx';d.save(path);return path
    def test_all_page_direct_style_numbering_header_fills_clear_without_changing_content(self):
        f=self.fixture();path=self.colored(f);before=content(path);evidence=scan_docx(path)
        self.assertTrue(evidence['issues']);self.assertTrue({'background','highlight','shd'}<={h['element'] for h in evidence['issues']})
        self.assertTrue({'word/document.xml','word/styles.xml','word/numbering.xml'}<={h['part'] for h in evidence['issues']})
        self.assertTrue(any('header' in h['part'] for h in evidence['issues']))
        result=normalize(path);self.assertFalse(result['evidence']['issues']);self.assertEqual(content(path),before)
        clean=sha(path);normalize(path);self.assertEqual(sha(path),clean)
    def test_generation_clears_imported_source_shading_and_preserves_source_hash(self):
        f=self.fixture();d=Document(f.source)
        for p in d.paragraphs:shade(p._p.get_or_add_pPr())
        shade(d.tables[0].cell(0,0)._tc.get_or_add_tcPr());d.save(f.source)
        original=sha(f.source);f.e=prepare(f.source,f.root/'evidence-colored')
        # Fixture build reads this path; keep fresh evidence bound to colored source.
        from bio_common import write
        write(f.root/'evidence/evidence.json',f.e)
        out=f.build();self.assertEqual(sha(f.source),original);self.assertFalse(scan_docx(out)['issues'])
        self.assertTrue(read(out.with_suffix('.trace.json'))['background']['removed'])
    def test_final_gate_rejects_color_reintroduced_after_generation(self):
        f=self.fixture();path=self.colored(f)
        summary=inspect_docx.run_inspection(path,f.root/'inspection','draft',True,True);report=read(summary['inspection'])
        self.assertTrue(report['biochemical_background_evidence']['issues'])
        with self.assertRaises(batch_tool.BatchError):bio_contract.check(report,f.e['title_candidate'])
    def test_gate_requires_current_complete_background_evidence(self):
        f=self.fixture();out=f.build();summary=inspect_docx.run_inspection(out,f.root/'inspection','draft',True,True);report=read(summary['inspection'])
        bio_contract.check(report,f.e['title_candidate'])
        for mutation in ('missing','stale','incomplete'):
            bad=copy.deepcopy(report)
            if mutation=='missing':bad.pop('biochemical_background_evidence')
            elif mutation=='stale':bad['biochemical_background_evidence']['artifact_sha256']='0'*64
            else:bad['biochemical_background_evidence']['complete']=False
            with self.assertRaises(batch_tool.BatchError):bio_contract.check(bad,f.e['title_candidate'])

if __name__=='__main__':unittest.main()
